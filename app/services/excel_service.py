"""
Genera el archivo Excel del "corte" de expedientes (botón "Generar corte",
pantallas Expedientes/Expedientes PA) -- una plantilla con encabezado,
filtros automáticos y columnas ajustadas, no un volcado plano de datos.
"""

from datetime import date, datetime
from enum import Enum
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_NAVY = "0B2942"
COLOR_BLUE_LIGHT = "DEECF9"
COLOR_WHITE = "FFFFFF"

# (encabezado, llave en el dict del expediente, ancho de columna)
COLUMNAS = [
    ("Expediente", "expediente", 14),
    ("Fecha Apertura", "fecha_apertura_servicio", 16),
    ("Cuenta", "cuenta", 26),
    ("Servicio", "tipo_servicio", 20),
    ("Subservicio", "tipo_subservicio", 20),
    ("Titular", "nombre_titular", 24),
    ("Paciente", "nombre_paciente", 24),
    ("Especialidad", "especialidad", 20),
    ("Entidad", "entidad", 16),
    ("Municipio", "municipio", 16),
    ("Fecha Asignación Proveedor", "fecha_asignacion_proveedor", 20),
    ("Fecha Cita", "fecha_cita", 14),
    ("Teléfono", "telefono", 14),
    ("Correo", "email", 26),
    ("Estado", "estatus", 20),
    # Evidencia de a qué hora se le mandó el expediente al proveedor (arranca
    # el conteo de 48h en Pago Anticipado) -- "NA" cuando el corte es de un
    # estatus distinto a "En Espera de Respuesta" (ver routers/expedientes.py).
    ("RFC Coordinador", "rfc_envio", 18),
    ("Fecha Envío a Proveedor", "fecha_envio", 22),
]

TITULOS_POR_TIPO = {
    "normal": "Corte de Expedientes",
    "anticipado": "Corte de Expedientes — Pago Anticipado",
}


def generar_corte_excel(registros: list[dict], tipo_expediente: str, usuario_nombre: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Corte"

    total_columnas = len(COLUMNAS)
    ultima_columna = get_column_letter(total_columnas)

    titulo = TITULOS_POR_TIPO.get(tipo_expediente, TITULOS_POR_TIPO["normal"])
    ws.merge_cells(f"A1:{ultima_columna}1")
    ws["A1"] = titulo
    ws["A1"].font = Font(size=16, bold=True, color=COLOR_WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=COLOR_NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(f"A2:{ultima_columna}2")
    generado = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws["A2"] = f"Generado por {usuario_nombre} el {generado} — {len(registros)} expediente(s)"
    ws["A2"].font = Font(size=10, italic=True, color=COLOR_NAVY)
    ws["A2"].fill = PatternFill("solid", fgColor=COLOR_BLUE_LIGHT)
    ws.row_dimensions[2].height = 18

    fila_encabezado = 4
    for col_idx, (encabezado, _clave, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=fila_encabezado, column=col_idx, value=encabezado)
        celda.font = Font(bold=True, color=COLOR_WHITE)
        celda.fill = PatternFill("solid", fgColor=COLOR_NAVY)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho
    ws.row_dimensions[fila_encabezado].height = 24

    for fila_idx, registro in enumerate(registros, start=fila_encabezado + 1):
        for col_idx, (_encabezado, clave, _ancho) in enumerate(COLUMNAS, start=1):
            valor = registro.get(clave)
            if isinstance(valor, Enum):
                # Ej. EstatusExpediente (str, Enum): sin esto, openpyxl escribe
                # "EstatusExpediente.EN_ESPERA_RESPUESTA" en vez del valor real.
                valor = valor.value
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            if isinstance(valor, datetime):
                celda.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(valor, date):
                celda.number_format = "yyyy-mm-dd"

    ultima_fila = fila_encabezado + len(registros)
    ws.auto_filter.ref = f"A{fila_encabezado}:{ultima_columna}{max(ultima_fila, fila_encabezado)}"
    ws.freeze_panes = f"A{fila_encabezado + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
