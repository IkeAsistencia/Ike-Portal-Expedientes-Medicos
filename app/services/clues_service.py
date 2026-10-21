"""
Enriquecimiento del padrón con el Catálogo CLUES de la Secretaría de Salud.

Para qué: el Excel del padrón trae las columnas de Latitud y Longitud, pero el
área las entrega vacías (662 de 662 filas), así que el mapa del Directorio no
tiene con qué ubicar a nadie. CLUES publica el domicilio Y las coordenadas de
los establecimientos de salud del país, y a diferencia de Google Maps no
restringe guardarlas (ver WEB-SCRAPING-DOCTORALIA.md, secciones 10 y 11).

De dónde sale el archivo: descarga directa de la DGIS, sin token ni scraping.
    http://www.dgis.salud.gob.mx/contenidos/intercambio/clues_gobmx.html
    -> "Establecimientos de Salud (Formato XLSX)"

Qué NO hace: CLUES no trae especialidad médica (el 100% de sus consultorios
tiene la subtipología en "NO ESPECIFICADO"), así que esto sirve para ubicar y
verificar proveedores que YA están en el padrón, no para darlos de alta.

EL CUIDADO IMPORTANTE -- por qué no basta con cruzar por RFC:
el RFC que trae CLUES es de la PERSONA MORAL, no de la sucursal. Medido contra
el archivo real: hay 159 RFC con más de un establecimiento (hasta 8 en un
caso), y cruzando solo por RFC, 12 proveedores del padrón empataban con un
establecimiento de OTRO municipio. Por eso todo cruce exige además que
coincidan municipio y código postal: es lo que distingue una sucursal de otra.
"""

import re
import unicodedata

from openpyxl import load_workbook

# La hoja lleva el corte en el nombre ("CLUES_202608"), que cambia en cada
# publicación: se busca por prefijo en vez de por nombre exacto.
PREFIJO_HOJA = "CLUES_"

INSTITUCION_PRIVADA = "SERVICIOS MEDICOS PRIVADOS"
ESTATUS_EN_OPERACION = "EN OPERACION"

LARGO_CODIGO_POSTAL = 5
# Un RFC de persona moral trae 12 caracteres y uno de persona física 13; menos
# que eso es un valor incompleto que no sirve para identificar a nadie.
LARGO_MINIMO_RFC = 12

# Columnas que se leen del catálogo (el archivo trae 68).
_COLUMNAS = {
    "clues": "CLUES",
    "institucion": "NOMBRE DE LA INSTITUCION",
    "estatus": "ESTATUS DE OPERACION",
    "nombre": "NOMBRE DE LA UNIDAD",
    "nombre_comercial": "NOMBRE COMERCIAL",
    "rfc": "RFC DEL ESTABLECIMIENTO",
    "estado": "ENTIDAD",
    "municipio": "MUNICIPIO",
    "codigo_postal": "CODIGO POSTAL",
    "tipologia": "NOMBRE DE TIPOLOGIA",
    "latitud": "LATITUD",
    "longitud": "LONGITUD",
}

# Con qué se reconoció al proveedor, de la más confiable a la menos. Se guarda
# en el reporte para poder revisar los cruces flojos antes de confiar en ellos.
CRUCE_RFC = "rfc+municipio+cp"
CRUCE_NOMBRE = "nombre+municipio+cp"


class ErrorCatalogoClues(Exception):
    """El archivo de CLUES no se pudo leer o no tiene el layout esperado."""


def normalizar(valor) -> str:
    """Mayúsculas, sin acentos y sin espacios de más."""
    base = str(valor or "").strip().upper()
    sin_acentos = "".join(
        c for c in unicodedata.normalize("NFD", base) if unicodedata.category(c) != "Mn"
    )
    return " ".join(sin_acentos.split())


def normalizar_rfc(valor) -> str:
    """Solo letras y dígitos: el padrón y CLUES los escriben distinto
    ("DIM-001129-HCA" vs "DIM001129HCA")."""
    return re.sub(r"[^A-Z0-9]", "", str(valor or "").upper())


def normalizar_cp(valor) -> str:
    """CP a 5 dígitos. Excel lo guarda como número y le come el cero de la
    izquierda, igual que pasaba en la carga del padrón."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    digitos = re.sub(r"\D", "", str(valor or ""))
    return digitos.zfill(LARGO_CODIGO_POSTAL) if digitos else ""


def _a_float(valor):
    if valor is None or str(valor).strip() == "":
        return None
    try:
        return float(str(valor).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None


def _hoja_de_establecimientos(wb):
    for nombre in wb.sheetnames:
        if nombre.upper().startswith(PREFIJO_HOJA):
            return wb[nombre]
    raise ErrorCatalogoClues(
        f'El archivo no trae la hoja de establecimientos (se esperaba una que empiece con "{PREFIJO_HOJA}"). '
        f"Hojas encontradas: {', '.join(wb.sheetnames) or 'ninguna'}."
    )


def leer_establecimientos(ruta, solo_privados: bool = True, solo_en_operacion: bool = True) -> list[dict]:
    """
    Los establecimientos del catálogo, ya filtrados y con los campos que
    interesan. Por defecto deja solo los privados en operación, que son los
    que pueden coincidir con el padrón del portal.
    """
    try:
        wb = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as e:
        raise ErrorCatalogoClues("El archivo de CLUES no es un Excel (.xlsx) válido.") from e

    ws = _hoja_de_establecimientos(wb)
    filas = ws.iter_rows(values_only=True)
    try:
        encabezado = [str(c or "").strip() for c in next(filas)]
    except StopIteration:
        raise ErrorCatalogoClues("La hoja de establecimientos está vacía.")

    indices = {clave: encabezado.index(col) for clave, col in _COLUMNAS.items() if col in encabezado}
    faltantes = [col for clave, col in _COLUMNAS.items() if clave not in indices]
    if faltantes:
        raise ErrorCatalogoClues(
            "Al archivo de CLUES le faltan columnas esperadas: " + ", ".join(faltantes) + "."
        )

    def celda(fila, clave):
        posicion = indices[clave]
        return fila[posicion] if posicion < len(fila) else None

    establecimientos = []
    for fila in filas:
        if all(v is None or str(v).strip() == "" for v in fila):
            continue
        if solo_privados and normalizar(celda(fila, "institucion")) != INSTITUCION_PRIVADA:
            continue
        if solo_en_operacion and normalizar(celda(fila, "estatus")) != ESTATUS_EN_OPERACION:
            continue
        establecimientos.append({
            "clues": str(celda(fila, "clues") or "").strip(),
            "nombre": str(celda(fila, "nombre") or "").strip(),
            "nombre_comercial": str(celda(fila, "nombre_comercial") or "").strip(),
            "rfc": normalizar_rfc(celda(fila, "rfc")),
            "estado": str(celda(fila, "estado") or "").strip(),
            "municipio": str(celda(fila, "municipio") or "").strip(),
            "codigo_postal": normalizar_cp(celda(fila, "codigo_postal")),
            "tipologia": str(celda(fila, "tipologia") or "").strip(),
            "latitud": _a_float(celda(fila, "latitud")),
            "longitud": _a_float(celda(fila, "longitud")),
        })
    wb.close()
    return establecimientos


def _zona(registro) -> str:
    """Municipio + CP: el desempate que evita asignarle a un proveedor las
    coordenadas de otra sucursal de la misma razón social."""
    return f"{normalizar(registro.get('municipio'))}|{normalizar_cp(registro.get('codigo_postal'))}"


def construir_indice(establecimientos) -> dict:
    """
    Índices de búsqueda, solo con los establecimientos que TRAEN coordenadas
    (sin ellas no sirven para enriquecer).

    Una llave que apunta a dos establecimientos distintos se descarta: es
    ambigua y preferimos no enriquecer a enriquecer mal.
    """
    por_rfc: dict[str, list] = {}
    por_nombre: dict[str, list] = {}
    for est in establecimientos:
        if est["latitud"] is None or est["longitud"] is None:
            continue
        zona = _zona(est)
        if len(est["rfc"]) >= LARGO_MINIMO_RFC:
            por_rfc.setdefault(f"{est['rfc']}|{zona}", []).append(est)
        # Un mismo establecimiento suele traer igual el nombre de la unidad y
        # el comercial: se indexa una sola vez por llave para que no se
        # confunda consigo mismo y quede marcado como ambiguo.
        for nombre in {normalizar(est["nombre"]), normalizar(est["nombre_comercial"])}:
            if nombre:
                por_nombre.setdefault(f"{nombre}|{zona}", []).append(est)

    def sin_ambiguos(indice):
        """Una llave que apunta a DOS establecimientos distintos se descarta.
        Se comparan por CLUES (su identificador) y no por objeto: el mismo
        establecimiento puede haber entrado por sus dos nombres."""
        limpio = {}
        for llave, candidatos in indice.items():
            distintos = {c["clues"]: c for c in candidatos}
            if len(distintos) == 1:
                limpio[llave] = next(iter(distintos.values()))
        return limpio

    return {"por_rfc": sin_ambiguos(por_rfc), "por_nombre": sin_ambiguos(por_nombre)}


def buscar_coordenadas(proveedor: dict, indice: dict):
    """
    (establecimiento, tipo_de_cruce) del CLUES que corresponde a ese proveedor,
    o (None, None) si no se le encontró uno confiable.

    Primero por RFC y luego por nombre, pero SIEMPRE exigiendo que coincidan
    municipio y código postal (ver el encabezado del módulo).
    """
    zona = _zona(proveedor)
    rfc = normalizar_rfc(proveedor.get("rfc"))
    if len(rfc) >= LARGO_MINIMO_RFC:
        est = indice["por_rfc"].get(f"{rfc}|{zona}")
        if est:
            return est, CRUCE_RFC

    nombre = normalizar(proveedor.get("nombre"))
    if nombre:
        est = indice["por_nombre"].get(f"{nombre}|{zona}")
        if est:
            return est, CRUCE_NOMBRE

    return None, None


def emparejar(proveedores, establecimientos) -> dict:
    """
    Qué coordenadas se le pondrían a cada proveedor del padrón.

    NO escribe nada: devuelve el plan para que quien llama decida (el script
    lo muestra antes de aplicarlo). Solo se proponen proveedores que hoy NO
    tienen coordenadas -- las que vinieron en el Excel del área son la fuente
    autorizada y no se tocan.
    """
    indice = construir_indice(establecimientos)
    con_coordenadas = sum(
        1 for e in establecimientos if e["latitud"] is not None and e["longitud"] is not None
    )
    asignaciones = []
    ya_tenian = 0
    sin_coincidencia = []

    for prov in proveedores:
        if prov.get("latitud") is not None and prov.get("longitud") is not None:
            ya_tenian += 1
            continue
        est, tipo = buscar_coordenadas(prov, indice)
        if est is None:
            sin_coincidencia.append(prov)
            continue
        asignaciones.append({
            "clave": prov.get("clave"),
            "nombre": prov.get("nombre"),
            "municipio": prov.get("municipio"),
            "latitud": est["latitud"],
            "longitud": est["longitud"],
            "cruce": tipo,
            "clues": est["clues"],
            "nombre_clues": est["nombre"],
        })

    return {
        "total_padron": len(proveedores),
        "ya_tenian_coordenadas": ya_tenian,
        "asignaciones": asignaciones,
        "sin_coincidencia": sin_coincidencia,
        "establecimientos_clues": len(establecimientos),
        "establecimientos_con_coordenadas": con_coordenadas,
    }
