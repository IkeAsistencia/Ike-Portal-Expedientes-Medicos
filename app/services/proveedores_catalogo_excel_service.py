"""
Carga del catálogo de proveedores del Directorio de Proveedores (pantalla
"Carga catálogo proveedores Core", exclusiva de Administrador) desde un
Excel -- la configuración y los datos de este catálogo viven SOLO en este
portal (SQLite local), nunca en SQL Server/Core. El archivo se solicita al
área correspondiente con el layout exacto que arma
`generar_plantilla_excel()`.
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_NAVY = "0B2942"
COLOR_WHITE = "FFFFFF"

# (encabezado en el Excel, llave interna, ancho de columna, obligatorio).
# cl_proveedor es la clave del proveedor en Core -- único campo realmente
# obligatorio además del nombre, para poder identificar a quién corresponde
# cada fila.
COLUMNAS = [
    ("Clave Proveedor (Core)", "cl_proveedor", 20, True),
    ("Nombre completo", "nombre", 30, True),
    ("Especialidad", "especialidad", 22, False),
    ("Cédula profesional", "cedula", 18, False),
    ("RFC", "rfc", 16, False),
    ("Estado", "estado", 18, False),
    ("Municipio", "municipio", 20, False),
    ("Dirección del consultorio", "direccion", 34, False),
    ("Nombre del consultorio/clínica", "consultorio", 26, False),
    ("Teléfono", "telefono", 16, False),
    ("Correo electrónico", "email", 28, False),
    ("Horario de atención", "horario", 24, False),
    ("Costo de consulta", "costo_consulta", 18, False),
    ("Convenio (aseguradora)", "convenio", 22, False),
    ("Teleconsulta (Sí/No)", "teleconsulta", 18, False),
]

FILA_ENCABEZADO = 1

_FILA_EJEMPLO = {
    "cl_proveedor": 10234, "nombre": "Dra. Laura Méndez Ruiz", "especialidad": "Cardiología",
    "cedula": "8453219", "rfc": "MERL850214H23", "estado": "Jalisco", "municipio": "Guadalajara",
    "direccion": "Av. Vallarta 1450, Col. Americana", "consultorio": "Consultorio 302",
    "telefono": "33 1234 5678", "email": "consultorio@correo.mx", "horario": "Lun-Vie 09:00-14:00",
    "costo_consulta": "800", "convenio": "GNP", "teleconsulta": "No",
}


def generar_plantilla_excel() -> bytes:
    """Layout a solicitar al área que entrega la data -- un encabezado por
    columna (los obligatorios marcados con "*") y un renglón de ejemplo
    para que quede claro el formato esperado de cada campo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Catálogo Proveedores"

    for col_idx, (encabezado, _clave, ancho, obligatorio) in enumerate(COLUMNAS, start=1):
        texto = encabezado + (" *" if obligatorio else "")
        celda = ws.cell(row=FILA_ENCABEZADO, column=col_idx, value=texto)
        celda.font = Font(bold=True, color=COLOR_WHITE)
        celda.fill = PatternFill("solid", fgColor=COLOR_NAVY)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho
    ws.row_dimensions[FILA_ENCABEZADO].height = 32
    ws.freeze_panes = f"A{FILA_ENCABEZADO + 1}"

    for col_idx, (_encabezado, clave, _ancho, _obligatorio) in enumerate(COLUMNAS, start=1):
        ws.cell(row=FILA_ENCABEZADO + 1, column=col_idx, value=_FILA_EJEMPLO.get(clave))

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ErrorCatalogoExcel(Exception):
    """El archivo subido no tiene el formato esperado -- encabezado
    faltante/renombrado, dato obligatorio vacío o de tipo inválido. El
    mensaje ya viene listo para mostrarse tal cual al usuario."""


def _encabezados_de_la_fila(fila: tuple) -> list[str]:
    encabezados = [str(c).strip() if c is not None else "" for c in fila]
    # La plantilla marca los obligatorios con " *" -- se compara sin ese
    # sufijo, por si alguien lo deja o lo borra al llenar el archivo.
    return [e[:-2].strip() if e.endswith(" *") else e for e in encabezados]


def leer_catalogo_excel(contenido: bytes) -> list[dict]:
    """Valida y convierte el Excel subido en una lista de dicts (una por
    proveedor), listos para `proveedores_catalogo_repo.reemplazar_catalogo`.
    Lanza ErrorCatalogoExcel con un mensaje accionable ante cualquier
    problema -- nunca intenta "adivinar" un dato faltante u obligatorio."""
    try:
        wb = load_workbook(BytesIO(contenido), data_only=True, read_only=True)
    except Exception as e:
        raise ErrorCatalogoExcel("El archivo no es un Excel (.xlsx) válido.") from e

    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise ErrorCatalogoExcel("El archivo está vacío.")

    encabezados = _encabezados_de_la_fila(filas[0])
    indice_por_clave: dict[str, int | None] = {}
    for encabezado_esperado, clave, _ancho, obligatorio in COLUMNAS:
        try:
            indice_por_clave[clave] = encabezados.index(encabezado_esperado)
        except ValueError:
            if obligatorio:
                raise ErrorCatalogoExcel(
                    f'Falta la columna obligatoria "{encabezado_esperado}" -- usa la plantilla '
                    "sin cambiarle los encabezados."
                )
            indice_por_clave[clave] = None

    registros = []
    for num_fila, fila in enumerate(filas[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in fila):
            continue  # renglón vacío (común al final de la plantilla) -- se ignora

        registro = {}
        for _encabezado, clave, _ancho, _obligatorio in COLUMNAS:
            idx = indice_por_clave[clave]
            valor = fila[idx] if idx is not None and idx < len(fila) else None
            if isinstance(valor, str):
                valor = valor.strip() or None
            registro[clave] = valor

        faltantes = [
            encabezado for encabezado, clave, _ancho, obligatorio in COLUMNAS
            if obligatorio and registro.get(clave) in (None, "")
        ]
        if faltantes:
            plural = "n" if len(faltantes) > 1 else ""
            raise ErrorCatalogoExcel(
                f"Fila {num_fila}: falta{plural} " + ", ".join(f'"{c}"' for c in faltantes) + "."
            )

        try:
            registro["cl_proveedor"] = int(registro["cl_proveedor"])
        except (TypeError, ValueError):
            raise ErrorCatalogoExcel(f'Fila {num_fila}: "Clave Proveedor (Core)" debe ser un número entero.')

        valor_tele = str(registro.get("teleconsulta") or "").strip().lower()
        registro["teleconsulta"] = 1 if valor_tele in ("si", "sí", "1", "true", "x") else 0

        registros.append(registro)

    if not registros:
        raise ErrorCatalogoExcel("El archivo no trae ningún renglón con datos.")

    return registros
