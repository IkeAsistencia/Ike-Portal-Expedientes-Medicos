"""
Lee un Excel de RFC/NOMBRE/ROL y da de alta esos accesos -- compartido por
`scripts/cargar_usuarios_excel.py` (uso manual, con modo seco) y por el
arranque de la app (siembra automática del primer despliegue, ver
`_sembrar_usuarios_iniciales` en app/main.py).

Formato esperado: ver el encabezado de scripts/cargar_usuarios_excel.py.
"""

import openpyxl

from app.repositories import accesos_repo
from app.schemas.acceso import PATRON_RFC, PATRON_RFC_ADMINISTRADOR

ROLES = {
    "CABINA": accesos_repo.PERFIL_CABINA,
    "ADMINISTRADOR": accesos_repo.PERFIL_ADMINISTRADOR,
    "PROVEEDOR": accesos_repo.PERFIL_PROVEEDOR,
}


class ErrorCargaUsuarios(Exception):
    pass


def leer_filas(ruta_archivo: str) -> list[dict]:
    try:
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
    except FileNotFoundError:
        raise ErrorCargaUsuarios(f"No existe el archivo: {ruta_archivo}")
    ws = wb.worksheets[0]
    encabezados = [
        (c.value or "").strip().upper() if isinstance(c.value, str) else c.value
        for c in next(ws.iter_rows(min_row=1, max_row=1))
    ]
    indice = {nombre: i for i, nombre in enumerate(encabezados) if nombre}

    for col in ("RFC", "NOMBRE", "ROL"):
        if col not in indice:
            raise ErrorCargaUsuarios(f"Al Excel le falta la columna {col!r} (encabezados encontrados: {encabezados}).")

    filas = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        rfc = fila[indice["RFC"]]
        if not rfc:
            continue
        filas.append({
            "rfc": str(rfc).strip().upper(),
            "nombre": str(fila[indice["NOMBRE"]] or "").strip(),
            "rol": str(fila[indice["ROL"]] or "").strip().upper(),
            "entidad": str(fila[indice["ENTIDAD"]]).strip() if "ENTIDAD" in indice and fila[indice["ENTIDAD"]] else None,
            "correo": str(fila[indice["CORREO"]]).strip() if "CORREO" in indice and fila[indice["CORREO"]] else None,
        })
    return filas


def validar_fila(fila: dict) -> str | None:
    """Regresa el motivo de rechazo, o None si la fila está lista para darse de alta."""
    if fila["rol"] not in ROLES:
        return f"rol desconocido {fila['rol']!r} (válidos: {', '.join(ROLES)})"
    if not fila["nombre"]:
        return "falta el nombre"

    perfil = ROLES[fila["rol"]]
    if perfil == accesos_repo.PERFIL_ADMINISTRADOR:
        if not PATRON_RFC_ADMINISTRADOR.match(fila["rfc"]):
            return "RFC de Administrador inválido: debe tener 11 caracteres (4 letras + 6 números + homoclave)"
    elif not PATRON_RFC.match(fila["rfc"]):
        return "RFC inválido: debe tener 10 caracteres (4 letras + 6 números)"

    if perfil == accesos_repo.PERFIL_PROVEEDOR and (not fila["entidad"] or not fila["correo"]):
        return "faltan ENTIDAD/CORREO, obligatorios para Proveedor"
    return None


def sembrar_desde_excel(ruta_archivo: str, aplicar: bool = True) -> dict:
    """
    Da de alta cada fila válida del Excel (alta_acceso es idempotente). Con
    aplicar=False solo valida y regresa el plan, sin escribir nada.
    """
    filas = leer_filas(ruta_archivo)
    listas, rechazadas = [], []
    for fila in filas:
        motivo = validar_fila(fila)
        (rechazadas if motivo else listas).append((fila, motivo))

    if aplicar:
        for fila, _ in listas:
            accesos_repo.alta_acceso(fila["rfc"], fila["nombre"], ROLES[fila["rol"]], fila["entidad"], fila["correo"])

    return {"total": len(filas), "listas": listas, "rechazadas": rechazadas}
