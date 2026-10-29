"""
Carga del padrón de proveedores del Directorio de Proveedores (pantalla
"Carga catálogo proveedores", exclusiva de Administrador) desde un Excel --
la configuración y los datos de este padrón viven SOLO en este portal
(SQLite local), nunca en SQL Server. El archivo se solicita al área
correspondiente con el layout exacto que arma `generar_plantilla_excel()`.

El archivo real (2026-10-28) trae DOS hojas, que se leen por separado y se
reconocen por cómo EMPIEZA su nombre (lo que venga después se ignora, p. ej.
"Padrón de proveedores 2026" cuenta igual que "Padrón de proveedores"):
  - "Padrón de proveedores": el padrón completo, origen 'catalogo'.
  - "Proveedores red de apoyo": subconjunto con menos columnas (sin cédula,
    RFC ni consultorio, y sin correo/horario/costo/convenio/teleconsulta/
    coordenadas -- ninguna de esas es obligatoria de por sí, ver
    `_registro_de_la_fila`), origen 'apoyo'. Se identifican aparte del
    padrón porque el área los entrega en una hoja distinta (instrucción del
    usuario, 2026-10-28) -- en el Directorio se ven con su propia etiqueta y
    color, no mezclados con "Dentro de la red".
Ambas hojas usan el mismo layout de columnas (ver COLUMNAS); basta con que
estén presentes las que cada hoja sí trae.

Layout vigente (2026-10-20), el que validó el área y que trae el archivo
real de 662 proveedores:
  - Los encabezados se toman TAL CUAL del archivo validado, asterisco
    incluido. OJO: ahí el "*" marca lo que NO es obligatorio, al revés de lo
    que significaba en la plantilla anterior.
  - NINGUNA columna se exige: una celda vacía entra como vacía en vez de
    rechazar el renglón (instrucción del usuario). El archivo real, por
    ejemplo, trae Correo, Latitud y Longitud vacías en las 662 filas.
  - Se quitó "Clave del proveedor": el portal genera la suya (PRV-000001,
    ver proveedores_repo) y es esa la que amarra las citas. Se sigue
    LEYENDO si el archivo la trae, para no romper los ya repartidos.
  - El domicilio va desglosado (calle con números, colonia y CP) y es el
    lugar donde se DA LA CONSULTA.
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.repositories.especialidades_repo import normalizar as _normalizar_texto
from app.repositories.proveedores_repo import ORIGEN_APOYO, ORIGEN_CATALOGO

COLOR_NAVY = "0B2942"
COLOR_WHITE = "FFFFFF"

# (encabezado en el Excel, llave interna, ancho de columna, otros nombres
# aceptados). El encabezado va EXACTO como en el archivo validado -- ya no
# se le pega el "*" por código, es parte del texto.
#
# Los "otros nombres" son los encabezados de plantillas anteriores: alguien
# puede traer todavía un archivo descargado antes de este cambio y no tiene
# por qué fallarle.
#
# Ninguna columna es obligatoria: ver _registro_de_la_fila.
COLUMNAS = [
    ("Nombre completo *", "nombre", 30, ("Nombre completo",)),
    ("Especialidad *", "especialidad", 22, ("Especialidad",)),
    ("Cédula profesional", "cedula", 18, ()),
    ("RFC *", "rfc", 16, ("RFC",)),
    ("Estado *", "estado", 18, ("Estado",)),
    ("Municipio *", "municipio", 20, ("Municipio",)),
    ("Calle, No. Ext., No. Int. *", "calle_num_ext_int", 38,
     ("Domicilio de consulta: Calle, No. Ext., No. Int.", "Calle, No. Ext., No. Int.")),
    ("Colonia *", "colonia", 24, ("Domicilio de consulta: Colonia", "Colonia")),
    ("Código Postal *", "codigo_postal", 20,
     ("Domicilio de consulta: Código Postal", "Código Postal")),
    ("Nombre del consultorio/clínica", "consultorio", 26, ()),
    ("Teléfono Consultorio", "telefono", 18, ("Teléfono",)),
    ("Correo electrónico*", "email", 28, ("Correo electrónico",)),
    ("Horario de atención", "horario", 24, ()),
    ("Costo de consulta", "costo_consulta", 18, ()),
    ("Convenio (aseguradora)", "convenio", 22, ()),
    ("Teleconsulta (Sí/No)", "teleconsulta", 18, ()),
    ("Latitud Consultorio*", "latitud", 18, ("Latitud",)),
    ("Longitud Consultorio*", "longitud", 18, ("Longitud",)),
]

# Ya no va en la plantilla, pero se sigue leyendo si el archivo la trae.
COLUMNA_CLAVE_CENTRAL = (
    "Clave del proveedor (referencia, opcional)", "cl_proveedor", 34,
    ("Clave del proveedor",),
)

FILA_ENCABEZADO = 1

# Con qué debe EMPEZAR el nombre de cada hoja reconocida (sin mayúsculas ni
# acentos, ver _normalizar_texto) -- lo que venga después se ignora.
_PREFIJO_HOJA_PADRON = _normalizar_texto("Padrón de proveedores")
_PREFIJO_HOJA_APOYO = _normalizar_texto("Proveedores red de apoyo")

_FILA_EJEMPLO = {
    "nombre": "Dra. Laura Méndez Ruiz", "especialidad": "CARDIOLOGIA",
    "cedula": "8453219", "rfc": "MERL850214H23", "estado": "Jalisco", "municipio": "Guadalajara",
    "calle_num_ext_int": "Av. Vallarta 1450, Int. 302", "colonia": "Americana",
    "codigo_postal": "44100", "consultorio": "Consultorio 302",
    "telefono": "33 1234 5678", "email": "consultorio@correo.mx", "horario": "Lun-Vie 09:00-14:00",
    "costo_consulta": "800", "convenio": "GNP", "teleconsulta": "No",
    "latitud": 20.6736, "longitud": -103.344,
}


def generar_plantilla_excel() -> bytes:
    """Layout a solicitar al área que entrega la data -- mismos encabezados
    que el archivo ya validado, más un renglón de ejemplo para que quede
    claro el formato esperado de cada campo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Padrón de Proveedores"

    for col_idx, (encabezado, _clave, ancho, _alias) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=FILA_ENCABEZADO, column=col_idx, value=encabezado)
        celda.font = Font(bold=True, color=COLOR_WHITE)
        celda.fill = PatternFill("solid", fgColor=COLOR_NAVY)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho
    ws.row_dimensions[FILA_ENCABEZADO].height = 44
    ws.freeze_panes = f"A{FILA_ENCABEZADO + 1}"

    for col_idx, (_encabezado, clave, _ancho, _alias) in enumerate(COLUMNAS, start=1):
        ws.cell(row=FILA_ENCABEZADO + 1, column=col_idx, value=_FILA_EJEMPLO.get(clave))

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ErrorCatalogoExcel(Exception):
    """El archivo subido no se puede leer: no es un Excel, viene vacío o no
    se le reconoció ni una columna. El mensaje ya viene listo para mostrarse
    tal cual al usuario.

    Un dato faltante NO llega aquí: entra vacío (ver _registro_de_la_fila)."""


def _llave_encabezado(texto) -> str:
    """
    Con qué se compara un encabezado: sin mayúsculas, sin acentos, sin
    asteriscos y sin espacios de más.

    Se ignora el asterisco a propósito -- en la plantilla vieja marcaba lo
    obligatorio y en la validada marca lo contrario, así que como dato no
    sirve, y además viene pegado o separado según la columna ("RFC *" pero
    "Correo electrónico*"). Comparar sin él evita que el archivo se rechace
    por un detalle de redacción.
    """
    return _normalizar_texto(str(texto or "").replace("*", " ")).lower()


def _indices_por_clave(fila_encabezado: tuple) -> dict[str, int]:
    """
    En qué columna del archivo quedó cada campo. Acepta el encabezado
    oficial o cualquiera de sus nombres anteriores, y no exige que vengan
    en orden ni que estén todos.
    """
    presentes = {}
    for posicion, celda in enumerate(fila_encabezado):
        llave = _llave_encabezado(celda)
        if llave and llave not in presentes:
            presentes[llave] = posicion

    indices = {}
    for encabezado, clave, _ancho, alias in list(COLUMNAS) + [COLUMNA_CLAVE_CENTRAL]:
        for nombre in (encabezado, *alias):
            posicion = presentes.get(_llave_encabezado(nombre))
            if posicion is not None:
                indices[clave] = posicion
                break
    return indices


LARGO_CODIGO_POSTAL = 5


def _codigo_postal(valor) -> str | None:
    """
    El CP como texto de 5 dígitos.

    Excel lo guarda como número y le come el cero de la izquierda: el
    archivo validado trae 166 filas así ("6700" en vez de "06700"). Se
    rellena a 5 (instrucción del usuario, 2026-10-20). Lo que no sea solo
    dígitos se deja tal cual -- puede ser un rango o una anotación.
    """
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = _texto(valor)
    if not texto:
        return None
    if texto.isdigit() and len(texto) < LARGO_CODIGO_POSTAL:
        return texto.zfill(LARGO_CODIGO_POSTAL)
    return texto


def _numero_o_none(valor) -> float | None:
    """Lat/long: si no es un número utilizable se deja en NULL en vez de
    rechazar el renglón (las 662 filas del archivo validado vienen sin
    geolocalización)."""
    texto = _texto(valor)
    if texto is None:
        return None
    try:
        return float(texto.replace(",", "."))
    except (TypeError, ValueError):
        return None


def _entero_o_none(valor) -> int | None:
    texto = _texto(valor)
    if texto is None:
        return None
    try:
        return int(float(texto))
    except (TypeError, ValueError):
        return None


# Texto que en realidad significa "no hay dato". El archivo validado trae 27
# celdas así ("NULL" en Estado, Municipio, Calle, Colonia, CP, Teléfono,
# Horario y Costo; "-" en RFC). Guardarlas tal cual se vería en la pantalla
# como "C.P. null" o "$NULL consulta", y además ensuciaría la llave de
# identidad del proveedor.
_TEXTO_VACIO = {"null", "nulo", "none", "n/a", "na", "s/d", "sin dato", "-", "--", "."}


def _texto(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto or texto.lower() in _TEXTO_VACIO:
        return None
    return texto


def _registro_de_la_fila(fila: tuple, indices: dict[str, int]) -> dict:
    """
    Un renglón del Excel convertido a los campos del padrón.

    NINGÚN campo se exige: lo que venga vacío queda en NULL y el proveedor se
    carga igual (instrucción del usuario, 2026-10-20). El nombre, que es lo
    único NOT NULL en la tabla, cae a cadena vacía -- así el renglón entra y
    se puede corregir después, en vez de perderse en la carga.
    """
    def celda(clave):
        posicion = indices.get(clave)
        if posicion is None or posicion >= len(fila):
            return None
        return fila[posicion]

    registro = {
        clave: _texto(celda(clave))
        for _encabezado, clave, _ancho, _alias in COLUMNAS
    }
    registro["nombre"] = registro.get("nombre") or ""
    registro["codigo_postal"] = _codigo_postal(celda("codigo_postal"))
    registro["latitud"] = _numero_o_none(celda("latitud"))
    registro["longitud"] = _numero_o_none(celda("longitud"))
    # Ya no viene en la plantilla; si el archivo la trae y no es un número,
    # se ignora en vez de tumbar la carga.
    registro["cl_proveedor"] = _entero_o_none(celda("cl_proveedor"))

    tele = _normalizar_texto(celda("teleconsulta")).lower()
    registro["teleconsulta"] = 1 if tele in ("si", "1", "true", "x", "verdadero") else 0
    return registro


def _hojas_del_catalogo(wb) -> list[tuple[str, object]]:
    """
    Las hojas reconocidas del archivo, con el origen que le toca a cada una
    (ver _PREFIJO_HOJA_PADRON / _PREFIJO_HOJA_APOYO). Si el archivo trae solo
    una de las dos, se carga igual con esa; si no reconoce ninguna, se lo
    deja a leer_catalogo_excel para que lo reporte con las hojas que sí trae.
    """
    hojas = []
    for nombre in wb.sheetnames:
        llave = _normalizar_texto(nombre)
        if llave.startswith(_PREFIJO_HOJA_PADRON):
            hojas.append((ORIGEN_CATALOGO, wb[nombre]))
        elif llave.startswith(_PREFIJO_HOJA_APOYO):
            hojas.append((ORIGEN_APOYO, wb[nombre]))
    return hojas


def leer_catalogo_excel(contenido: bytes) -> list[dict]:
    """
    Convierte el Excel subido en una lista de dicts (uno por proveedor),
    listos para `proveedores_repo.cargar_padron` -- cada uno con su "origen"
    ('catalogo' o 'apoyo') según de cuál de las dos hojas salió.

    Solo lanza ErrorCatalogoExcel si el archivo no se puede leer de plano (no
    es Excel, no trae ninguna hoja reconocida, o ninguna de las reconocidas
    tiene un encabezado utilizable). Un dato faltante dentro de un renglón NO
    detiene la carga: entra vacío.
    """
    try:
        wb = load_workbook(BytesIO(contenido), data_only=True, read_only=True)
    except Exception as e:
        raise ErrorCatalogoExcel("El archivo no es un Excel (.xlsx) válido.") from e

    hojas = _hojas_del_catalogo(wb)
    if not hojas:
        raise ErrorCatalogoExcel(
            'El archivo no trae ninguna hoja reconocida -- se esperaba una que empiece con '
            '"Padrón de proveedores" o "Proveedores red de apoyo". '
            f"Hojas encontradas: {', '.join(wb.sheetnames) or 'ninguna'}."
        )

    registros = []
    for origen, ws in hojas:
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue  # hoja reconocida pero vacía -- se ignora, no tumba la carga completa

        indices = _indices_por_clave(filas[0])
        if not indices:
            continue  # encabezado irreconocible en esta hoja -- idem

        for fila in filas[1:]:
            if all(c is None or str(c).strip() == "" for c in fila):
                continue  # renglón vacío (común al final de la plantilla) -- se ignora
            registro = _registro_de_la_fila(fila, indices)
            registro["origen"] = origen
            registros.append(registro)

    if not registros:
        raise ErrorCatalogoExcel(
            "No se reconoció ninguna columna del encabezado en ninguna hoja, o el archivo no "
            "trae ningún renglón con datos -- descarga la plantilla y úsala sin cambiarle los "
            "nombres a las columnas."
        )

    return registros
