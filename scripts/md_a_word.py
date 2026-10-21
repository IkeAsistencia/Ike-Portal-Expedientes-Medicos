"""
Convierte un documento Markdown del proyecto a Word (.docx), para compartirlo
con áreas que no leen el repositorio (Jurídico, Negocio, Compras).

Se escribió a la medida de los documentos de análisis de este proyecto
(WEB-SCRAPING-DOCTORALIA.md, SECURITY-REVIEW.md): maneja encabezados, listas,
tablas, citas en bloque, bloques de código y negritas/código en línea. No
pretende ser un convertidor general de Markdown.

USO:
    python -m scripts.md_a_word WEB-SCRAPING-DOCTORALIA.md
    python -m scripts.md_a_word WEB-SCRAPING-DOCTORALIA.md --salida "Analisis.docx"

El .docx NO se versiona (los binarios van en .gitignore): se regenera cuando
haga falta, así el Markdown sigue siendo la única fuente de verdad.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

# Markdown en línea que sí se interpreta: **negritas** y `código`.
_TROZOS = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
_LIGA = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
GRIS = RGBColor(0x44, 0x44, 0x44)


def _texto_plano(texto: str) -> str:
    """Quita la sintaxis que no se va a representar como formato."""
    # Las ligas quedan como "texto (url)": en papel la URL hay que poder leerla.
    texto = _LIGA.sub(lambda m: f"{m.group(1)} ({m.group(2)})", texto)
    return texto.replace("\\", "")


def _escribir_con_formato(parrafo, texto: str) -> None:
    """Agrega el texto al párrafo respetando negritas y código en línea."""
    for trozo in _TROZOS.split(_texto_plano(texto)):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            run = parrafo.add_run(trozo[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            parrafo.add_run(trozo)


def _es_separador_de_tabla(linea: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", linea.strip()))


def _celdas(linea: str) -> list[str]:
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def _agregar_tabla(doc, filas: list[list[str]]) -> None:
    if not filas:
        return
    columnas = max(len(f) for f in filas)
    tabla = doc.add_table(rows=0, cols=columnas)
    tabla.style = "Light Grid Accent 1"
    for i, fila in enumerate(filas):
        celdas = tabla.add_row().cells
        for j in range(columnas):
            texto = fila[j] if j < len(fila) else ""
            parrafo = celdas[j].paragraphs[0]
            _escribir_con_formato(parrafo, texto)
            if i == 0:
                for run in parrafo.runs:
                    run.bold = True


def _agregar_codigo(doc, lineas: list[str]) -> None:
    parrafo = doc.add_paragraph()
    parrafo.paragraph_format.left_indent = Pt(18)
    parrafo.paragraph_format.space_after = Pt(10)
    run = parrafo.add_run("\n".join(lineas))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = GRIS


def convertir(ruta_md: Path, ruta_docx: Path) -> Path:
    lineas = ruta_md.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    i = 0
    tabla_pendiente: list[list[str]] = []

    def cerrar_tabla():
        nonlocal tabla_pendiente
        if tabla_pendiente:
            _agregar_tabla(doc, tabla_pendiente)
            doc.add_paragraph()
            tabla_pendiente = []

    while i < len(lineas):
        linea = lineas[i]
        despojada = linea.strip()

        # Bloque de código
        if despojada.startswith("```"):
            cerrar_tabla()
            i += 1
            bloque = []
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i])
                i += 1
            _agregar_codigo(doc, bloque)
            i += 1
            continue

        # Tabla (se acumula y se escribe al terminar)
        if despojada.startswith("|"):
            if not _es_separador_de_tabla(despojada):
                tabla_pendiente.append(_celdas(despojada))
            i += 1
            continue
        cerrar_tabla()

        if not despojada:
            i += 1
            continue

        # Encabezados
        encabezado = re.match(r"^(#{1,6})\s+(.*)$", despojada)
        if encabezado:
            nivel = len(encabezado.group(1))
            titulo = _texto_plano(encabezado.group(2)).replace("*", "").replace("`", "")
            doc.add_heading(titulo, level=min(nivel, 4))
            i += 1
            continue

        # Línea horizontal
        if re.fullmatch(r"[-*_]{3,}", despojada):
            doc.add_paragraph().add_run("―" * 30).font.color.rgb = GRIS
            i += 1
            continue

        # Cita en bloque (se juntan los renglones continuos en un solo párrafo)
        if despojada.startswith(">"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                bloque.append(lineas[i].strip().lstrip(">").strip())
                i += 1
            parrafo = doc.add_paragraph()
            parrafo.paragraph_format.left_indent = Pt(24)
            parrafo.paragraph_format.space_after = Pt(10)
            parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _escribir_con_formato(parrafo, " ".join(bloque))
            for run in parrafo.runs:
                run.italic = True
            continue

        # Listas (numeradas y con viñeta); se pegan los renglones de
        # continuación, que en el Markdown vienen indentados.
        vineta = re.match(r"^[-*]\s+(.*)$", despojada)
        numerada = re.match(r"^(\d+)[.)]\s+(.*)$", despojada)
        if vineta or numerada:
            texto = (vineta or numerada).group(1 if vineta else 2)
            i += 1
            while (
                i < len(lineas)
                and lineas[i].strip()
                and lineas[i].startswith(("   ", "\t"))
                and not re.match(r"^\s*([-*]|\d+[.)])\s", lineas[i])
                and not lineas[i].strip().startswith(("|", "```"))
            ):
                texto += " " + lineas[i].strip()
                i += 1
            estilo = "List Number" if numerada else "List Bullet"
            _escribir_con_formato(doc.add_paragraph(style=estilo), texto)
            continue

        # Párrafo normal: se juntan los renglones hasta la línea en blanco
        # (el Markdown viene cortado a 80 columnas).
        bloque = [despojada]
        i += 1
        while (
            i < len(lineas)
            and lineas[i].strip()
            and not re.match(r"^\s*(#{1,6}\s|[-*]\s|\d+[.)]\s|>|\||```|[-*_]{3,}$)", lineas[i])
        ):
            bloque.append(lineas[i].strip())
            i += 1
        parrafo = doc.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _escribir_con_formato(parrafo, " ".join(bloque))

    cerrar_tabla()
    doc.save(ruta_docx)
    return ruta_docx


def main() -> None:
    parser = argparse.ArgumentParser(description="Convierte un .md del proyecto a .docx")
    parser.add_argument("archivo", help="Ruta del Markdown")
    parser.add_argument("--salida", help="Ruta del .docx (por defecto, el mismo nombre)")
    args = parser.parse_args()

    ruta_md = Path(args.archivo)
    if not ruta_md.exists():
        raise SystemExit(f"No se encontró el archivo: {ruta_md}")
    ruta_docx = Path(args.salida) if args.salida else ruta_md.with_suffix(".docx")

    convertir(ruta_md, ruta_docx)
    print(f"Generado: {ruta_docx}  ({ruta_docx.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
