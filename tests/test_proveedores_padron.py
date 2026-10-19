"""
Pruebas del padrón de proveedores del Directorio de Proveedores:
/proveedores-catalogo (la carga del Excel, exclusiva de Administrador) y
/directorio-proveedores (lectura y citas, abierto a Proveedor y
Administrador) -- NO requiere SQL Server real: todo vive en SQLite local.
"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services.proveedores_excel_service import COLUMNAS, generar_plantilla_excel

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _sin_sql_server(monkeypatch):
    import app.main as main_module

    monkeypatch.setattr(main_module, "calentar_pool", lambda: None)

    # El limitador de intentos de /auth (app/core/rate_limit.py) guarda su
    # conteo en memoria del PROCESO, compartido entre archivos de prueba --
    # sin esto, los mismos RFC de prueba (RFCADMTEST, etc.) usados también
    # en test_smoke.py chocan entre sí y el login empieza a responder 429.
    from app.core.rate_limit import limpiar_todo

    limpiar_todo()


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_AUTH_MODE", "sql")
    monkeypatch.setenv("DB_SERVER", "servidor-fake")
    monkeypatch.setenv("DB_NAME", "bd-fake")
    monkeypatch.setenv("DB_USER", "usuario-fake")
    monkeypatch.setenv("DB_PASSWORD", "password-fake")
    monkeypatch.setenv("LOCAL_DB_PATH", str(tmp_path / "test_local.db"))
    monkeypatch.setenv("JWT_SECRET_KEY", "clave-de-prueba")

    from app.config import get_settings

    get_settings.cache_clear()

    with TestClient(app) as c:
        yield c


def _auth_headers_admin(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCADMTEST", "Admin de Prueba", accesos_repo_module.PERFIL_ADMINISTRADOR)
    accesos_repo_module.crear_password("RFCADMTEST", "Admin2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCADMTEST", "password": "Admin2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_proveedor(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCPROTEST", "Proveedor de Prueba", accesos_repo_module.PERFIL_PROVEEDOR)
    accesos_repo_module.crear_password("RFCPROTEST", "Proveedor2024")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCPROTEST", "password": "Proveedor2024"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_cabina(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCCABTEST", "Cabina de Prueba", accesos_repo_module.PERFIL_CABINA)
    accesos_repo_module.crear_password("RFCCABTEST", "Cabina2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCCABTEST", "password": "Cabina2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _xlsx_de_filas(filas: list[dict]) -> bytes:
    """Arma un .xlsx con los mismos encabezados de la plantilla real, para
    no depender de generar_plantilla_excel() salvo en la prueba que la
    prueba explícitamente a ella."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col_idx, (encabezado, _clave, _ancho, obligatorio) in enumerate(COLUMNAS, start=1):
        ws.cell(row=1, column=col_idx, value=encabezado + (" *" if obligatorio else ""))
    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, (_encabezado, clave, _ancho, _obligatorio) in enumerate(COLUMNAS, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=fila.get(clave))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _cargar(client, headers, filas):
    return client.post(
        "/proveedores-catalogo/cargar",
        headers=headers,
        files={"archivo": ("catalogo.xlsx", _xlsx_de_filas(filas), XLSX_MIME)},
    )


FILA_VALIDA = {
    "nombre": "Dra. Laura Méndez Ruiz", "especialidad": "Cardiología",
    "cedula": "8453219", "rfc": "MERL850214H23", "estado": "Jalisco", "municipio": "Guadalajara",
    "calle_num_ext_int": "Av. Vallarta 1450, Int. 302", "colonia": "Americana",
    "codigo_postal": "44100", "consultorio": "Consultorio 302", "telefono": "33 1234 5678",
    "email": "consultorio@correo.mx", "horario": "Lun-Vie 09:00-14:00", "costo_consulta": "800",
    "convenio": "GNP", "teleconsulta": "No",
}


# ---------- Acceso exclusivo de Administrador ----------

def test_plantilla_requiere_administrador(client):
    r = client.get("/proveedores-catalogo/plantilla", headers=_auth_headers_proveedor(client))
    assert r.status_code == 403


def test_cargar_requiere_administrador(client):
    r = _cargar(client, _auth_headers_cabina(client), [FILA_VALIDA])
    assert r.status_code == 403


def test_resumen_requiere_administrador(client):
    r = client.get("/proveedores-catalogo", headers=_auth_headers_proveedor(client))
    assert r.status_code == 403


# ---------- Plantilla ----------

def test_plantilla_se_descarga_y_abre_como_excel_valido(client):
    r = client.get("/proveedores-catalogo/plantilla", headers=_auth_headers_admin(client))
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX_MIME
    wb = load_workbook(BytesIO(r.content))
    encabezados = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    # El nombre es ahora el único obligatorio: la clave del proveedor dejó
    # de serlo (decisión del usuario, 2026-10-19).
    assert encabezados[0] == "Nombre completo *"
    assert [e for e in encabezados if e.endswith(" *")] == ["Nombre completo *"]


def test_la_plantilla_trae_el_domicilio_de_consulta_desglosado(client):
    """La dirección va separada y etiquetada como el domicilio DONDE DA LA
    CONSULTA -- ya no es un solo campo de texto libre."""
    r = client.get("/proveedores-catalogo/plantilla", headers=_auth_headers_admin(client))
    encabezados = [c.value for c in next(load_workbook(BytesIO(r.content)).active.iter_rows(min_row=1, max_row=1))]
    assert "Domicilio de consulta: Calle, No. Ext., No. Int." in encabezados
    assert "Domicilio de consulta: Colonia" in encabezados
    assert "Domicilio de consulta: Código Postal" in encabezados
    assert not any("Dirección del consultorio" == e for e in encabezados)


def test_la_plantilla_generada_se_puede_leer_de_vuelta(client):
    """La plantilla que se descarga (con su fila de ejemplo) debe pasar la
    misma validación que un archivo real -- si no, el ejemplo que le damos
    al área estaría mal."""
    from app.services.proveedores_excel_service import leer_catalogo_excel

    registros = leer_catalogo_excel(generar_plantilla_excel())
    assert len(registros) == 1
    assert registros[0]["calle_num_ext_int"] == "Av. Vallarta 1450, Int. 302"
    assert registros[0]["colonia"] == "Americana"
    assert registros[0]["codigo_postal"] == "44100"


# ---------- Carga ----------

def test_cargar_padron_lo_deja_disponible_en_el_resumen(client):
    headers = _auth_headers_admin(client)
    r = _cargar(client, headers, [
        FILA_VALIDA,
        {**FILA_VALIDA, "nombre": "Dr. Otro Médico", "cedula": "5555555"},
    ])
    assert r.status_code == 200, r.text
    assert r.json() == {"ok": True, "total": 2, "nuevos": 2, "actualizados": 0, "absorbidos": 0}

    resumen = client.get("/proveedores-catalogo", headers=headers).json()
    assert resumen["total"] == 2
    assert resumen["dentro_de_la_red"] == 2
    assert resumen["ultima_carga"]["rfc"] == "RFCADMTEST"
    assert resumen["ultima_carga"]["nuevos"] == 2


def test_cada_proveedor_recibe_una_clave_propia_del_portal(client):
    """La clave la genera el portal (PRV-000001): es la que amarra las citas
    y ya no depende de que el Excel traiga la del sistema central."""
    headers = _auth_headers_admin(client)
    _cargar(client, headers, [FILA_VALIDA, {**FILA_VALIDA, "nombre": "Dr. Dos", "cedula": "111111"}])

    padron = client.get("/directorio-proveedores", headers=headers).json()
    claves = sorted(p["clave"] for p in padron)
    assert claves == ["PRV-000001", "PRV-000002"]
    assert all(p["origen"] == "catalogo" for p in padron)


def test_se_puede_cargar_sin_clave_de_proveedor(client):
    headers = _auth_headers_admin(client)
    r = _cargar(client, headers, [{**FILA_VALIDA, "cl_proveedor": None}])
    assert r.status_code == 200, r.text

    padron = client.get("/directorio-proveedores", headers=headers).json()
    assert padron[0]["cl_proveedor"] is None
    assert padron[0]["clave"] == "PRV-000001"


def test_segunda_carga_agrega_los_nuevos_y_conserva_la_clave_de_los_que_ya_estaban(client):
    """Decisión del usuario (2026-10-19): la carga ya NO reemplaza el padrón.
    Agrega lo que no existe y actualiza lo que sí, SIN cambiarle la clave --
    si cambiara, las citas ya agendadas apuntarían a otro médico."""
    headers = _auth_headers_admin(client)
    _cargar(client, headers, [FILA_VALIDA])
    clave_original = client.get("/directorio-proveedores", headers=headers).json()[0]["clave"]

    r = _cargar(client, headers, [
        {**FILA_VALIDA, "costo_consulta": "950"},           # el mismo, con otro precio
        {**FILA_VALIDA, "nombre": "Dr. Nuevo", "cedula": "7777777"},
    ])
    assert r.json() == {"ok": True, "total": 2, "nuevos": 1, "actualizados": 1, "absorbidos": 0}

    padron = {p["nombre"]: p for p in client.get("/directorio-proveedores", headers=headers).json()}
    assert len(padron) == 2
    assert padron[FILA_VALIDA["nombre"]]["clave"] == clave_original
    assert padron[FILA_VALIDA["nombre"]]["costo_consulta"] == "950"


def test_un_medico_con_varias_especialidades_ocupa_un_renglon_por_especialidad(client):
    """Por eso la identidad no puede ser solo la cédula: el mismo médico con
    dos especialidades son dos renglones, cada uno con su clave y su agenda."""
    headers = _auth_headers_admin(client)
    r = _cargar(client, headers, [
        FILA_VALIDA,
        {**FILA_VALIDA, "especialidad": "Medicina General"},
    ])
    assert r.json()["nuevos"] == 2

    padron = client.get("/directorio-proveedores", headers=headers).json()
    assert len({p["clave"] for p in padron}) == 2
    assert {p["especialidad"] for p in padron} == {"Cardiología", "Medicina General"}


def test_recargar_el_mismo_archivo_no_duplica_ni_multiplica_claves(client):
    headers = _auth_headers_admin(client)
    _cargar(client, headers, [FILA_VALIDA])
    _cargar(client, headers, [FILA_VALIDA])
    r = _cargar(client, headers, [FILA_VALIDA])
    assert r.json()["nuevos"] == 0
    assert r.json()["actualizados"] == 1
    assert len(client.get("/directorio-proveedores", headers=headers).json()) == 1


# ---------- Validaciones del archivo ----------

def test_rechaza_archivo_que_no_es_xlsx(client):
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.txt", b"no es un excel", "text/plain")},
    )
    assert r.status_code == 400


def test_rechaza_xlsx_falso_con_content_type_mentiroso(client):
    """El Content-Type lo manda el cliente -- se valida también la firma
    real del archivo (un .xlsx es un ZIP, empieza con PK\\x03\\x04)."""
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.xlsx", b"esto no es un zip valido", XLSX_MIME)},
    )
    assert r.status_code == 400


def test_rechaza_archivo_sin_la_columna_obligatoria(client):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Especialidad")  # falta "Nombre completo *"
    ws.cell(row=2, column=1, value="Cardiología")
    buffer = BytesIO()
    wb.save(buffer)

    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.xlsx", buffer.getvalue(), XLSX_MIME)},
    )
    assert r.status_code == 400
    assert "Nombre completo" in r.json()["detail"]


def test_rechaza_fila_sin_nombre(client):
    r = _cargar(client, _auth_headers_admin(client), [{**FILA_VALIDA, "nombre": None}])
    assert r.status_code == 400
    assert "Nombre completo" in r.json()["detail"]


def test_rechaza_clave_de_proveedor_no_numerica(client):
    r = _cargar(client, _auth_headers_admin(client), [{**FILA_VALIDA, "cl_proveedor": "ABC"}])
    assert r.status_code == 400
    assert "número entero" in r.json()["detail"]


def test_el_codigo_postal_conserva_los_ceros_a_la_izquierda(client):
    """Excel suele mandar el CP como número; guardarlo así se comería el cero
    de "01000"."""
    headers = _auth_headers_admin(client)
    _cargar(client, headers, [{**FILA_VALIDA, "codigo_postal": "01000"}])
    padron = client.get("/directorio-proveedores", headers=headers).json()
    assert padron[0]["codigo_postal"] == "01000"


def test_ignora_renglones_vacios_al_final(client):
    r = _cargar(client, _auth_headers_admin(client), [FILA_VALIDA, {}, {}])
    assert r.status_code == 200
    assert r.json()["total"] == 1


def test_teleconsulta_si_no_se_interpreta_a_booleano(client):
    headers = _auth_headers_admin(client)
    _cargar(client, headers, [
        {**FILA_VALIDA, "cedula": "1", "teleconsulta": "Sí"},
        {**FILA_VALIDA, "cedula": "2", "teleconsulta": "No"},
    ])
    padron = {p["cedula"]: p for p in client.get("/directorio-proveedores", headers=headers).json()}
    assert padron["1"]["teleconsulta"] == 1
    assert padron["2"]["teleconsulta"] == 0


# ---------- GET /directorio-proveedores (lectura para la pantalla "Proveedores") ----------

def test_directorio_rechaza_cabina(client):
    r = client.get("/directorio-proveedores", headers=_auth_headers_cabina(client))
    assert r.status_code == 403


def test_directorio_permite_proveedor_y_administrador(client):
    assert client.get("/directorio-proveedores", headers=_auth_headers_proveedor(client)).status_code == 200
    assert client.get("/directorio-proveedores", headers=_auth_headers_admin(client)).status_code == 200
