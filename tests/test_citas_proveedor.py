"""
Pruebas de las citas del Directorio de Proveedores
(/directorio-proveedores/citas y /{clave}/agenda) -- NO requiere SQL Server
real: todo vive en SQLite local y el expediente se simula.

Antes de esto, "Agendar cita" era una simulación del frontend: los horarios
se inventaban con un hash y nada quedaba registrado. Ahora la cita se guarda
contra la clave del proveedor (PRV-000001), que es la que no cambia entre
cargas del Excel, y SIEMPRE para un expediente (2026-10-26): antes una cita
agendada desde el menú quedaba sin paciente ni expediente.
"""
from datetime import date, datetime, time, timedelta
from io import BytesIO
from itertools import count

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services import agenda_proveedores_service
from app.services.agenda_proveedores_service import DIAS_CITA_MAXIMO, HORARIOS_BASE
from app.services.proveedores_excel_service import COLUMNAS

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _sin_sql_server(monkeypatch):
    import app.main as main_module

    monkeypatch.setattr(main_module, "calentar_pool", lambda: None)

    from app.core.rate_limit import limpiar_todo

    limpiar_todo()


@pytest.fixture(autouse=True)
def _reloj_fijo(monkeypatch):
    """Las pruebas agendan "para hoy": con el reloj real fallarían según la hora
    a la que se corran (una hora de hoy que ya pasó se rechaza). Se fija a las
    06:00 de hoy; las pruebas de "hora pasada" lo cambian a lo que necesiten."""
    monkeypatch.setattr(
        agenda_proveedores_service, "ahora", lambda: datetime.combine(date.today(), time(6, 0))
    )


def _falso_expediente(monkeypatch, estatus="En Espera de Respuesta", existe=True):
    """El expediente viene de SQL Server en producción; aquí se simula. Cualquier
    número de expediente existe con el estatus indicado."""
    import app.repositories.expedientes_repo as expedientes_repo
    from app.schemas.expediente import EstatusExpediente

    def listar(filtro):
        if not existe:
            return []
        return [_expediente(filtro.cl_expediente, EstatusExpediente(estatus))]

    monkeypatch.setattr(expedientes_repo, "listar_expedientes", listar)


def _expediente(cl, estatus):
    from app.schemas.expediente import Expediente

    return Expediente(
        expediente=cl, cuenta="Cuenta Demo", tipo_servicio="Asistencia Médica",
        tipo_subservicio="Consulta Médica en Consultorio", nombre_paciente="Juan de Dios Peralta",
        especialidad="CARDIOLOGIA", entidad="JALISCO", municipio="GUADALAJARA",
        fecha_apertura_servicio=date.today(), fecha_apertura_completa=date.today().isoformat() + "T08:00:00",
        estatus=estatus,
    )


@pytest.fixture(autouse=True)
def _expediente_en_espera(monkeypatch):
    _falso_expediente(monkeypatch)


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_AUTH_MODE", "sql")
    monkeypatch.setenv("DB_SERVER", "servidor-fake")
    monkeypatch.setenv("DB_NAME", "bd-fake")
    monkeypatch.setenv("DB_USER", "usuario-fake")
    monkeypatch.setenv("DB_PASSWORD", "password-fake")
    monkeypatch.setenv("LOCAL_DB_PATH", str(tmp_path / "test_local.db"))
    monkeypatch.setenv("JWT_SECRET_KEY", "clave-de-prueba")

    from app.config import get_settings

    get_settings.cache_clear()

    with TestClient(app) as c:
        yield c


def _auth_headers_admin(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCADMTEST", "Admin de Prueba", accesos_repo_module.PERFIL_ADMINISTRADOR)
    accesos_repo_module.crear_password("RFCADMTEST", "Admin2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCADMTEST", "password": "Admin2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_cabina(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCCABTEST", "Cabina de Prueba", accesos_repo_module.PERFIL_CABINA)
    accesos_repo_module.crear_password("RFCCABTEST", "Cabina2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCCABTEST", "password": "Cabina2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_proveedor(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso(
        "RFCPROTEST", "Proveedor de Prueba", accesos_repo_module.PERFIL_PROVEEDOR, "Jalisco", "prov@ejemplo.com"
    )
    accesos_repo_module.crear_password("RFCPROTEST", "Proveedor2024")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCPROTEST", "password": "Proveedor2024"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _cargar_proveedores(client, headers, filas):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col_idx, (encabezado, _clave, _ancho, _alias) in enumerate(COLUMNAS, start=1):
        ws.cell(row=1, column=col_idx, value=encabezado)
    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, (_encabezado, clave, _ancho, _alias) in enumerate(COLUMNAS, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=fila.get(clave))
    buffer = BytesIO()
    wb.save(buffer)
    client.post(
        "/proveedores-catalogo/cargar", headers=headers,
        files={"archivo": ("catalogo.xlsx", buffer.getvalue(), XLSX_MIME)},
    )
    return client.get("/directorio-proveedores", headers=headers).json()


@pytest.fixture()
def proveedor(client):
    """Un proveedor ya cargado, con su clave del portal lista para agendar.
    SIN horario de atención: así se ofrecen los horarios genéricos y las
    pruebas no dependen del día de la semana en que se corran (el horario
    tiene sus propias pruebas más abajo)."""
    headers = _auth_headers_admin(client)
    padron = _cargar_proveedores(client, headers, [{
        "nombre": "Dra. Laura Méndez Ruiz", "especialidad": "Cardiología", "cedula": "8453219",
        "estado": "Jalisco", "municipio": "Guadalajara",
    }])
    return {"headers": headers, "clave": padron[0]["clave"]}


@pytest.fixture()
def prov_headers(client):
    return _auth_headers_proveedor(client)


HOY = date.today().isoformat()
_numeros = count(7001)


def _agendar(client, headers, proveedor, hora="09:00", fecha=None, expediente=None):
    """Cada llamada usa un expediente distinto, salvo que se indique uno: un
    expediente solo puede tener una cita activa."""
    return client.post("/directorio-proveedores/citas", headers=headers, json={
        "clave_proveedor": proveedor["clave"], "fecha": fecha or HOY, "hora": hora,
        "cl_expediente": expediente or next(_numeros),
    })


def _citas(client, headers):
    return client.get("/directorio-proveedores/citas", headers=headers).json()


# ---------- Agendar ----------

def test_agendar_cita_queda_registrada_con_folio_y_el_paciente_del_expediente(client, proveedor):
    r = _agendar(client, proveedor["headers"], proveedor, expediente=5001)
    assert r.status_code == 200, r.text
    assert r.json()["folio"].startswith("CITA-")

    citas = _citas(client, proveedor["headers"])
    assert len(citas) == 1
    assert citas[0]["clave_proveedor"] == proveedor["clave"]
    assert citas[0]["hora"] == "09:00"
    assert citas[0]["paciente"] == "Juan de Dios Peralta"
    assert citas[0]["cl_expediente"] == 5001
    assert citas[0]["estatus"] == "Agendada"
    assert citas[0]["rfc_agendo"] == "RFCADMTEST"


def test_toda_cita_es_para_un_expediente(client, proveedor):
    r = client.post("/directorio-proveedores/citas", headers=proveedor["headers"], json={
        "clave_proveedor": proveedor["clave"], "fecha": HOY, "hora": "13:30", "paciente": "Ana",
    })
    assert r.status_code == 422
    assert _citas(client, proveedor["headers"]) == []


def test_no_se_puede_apartar_dos_veces_el_mismo_horario(client, proveedor):
    assert _agendar(client, proveedor["headers"], proveedor, hora="10:30").status_code == 200
    repetida = _agendar(client, proveedor["headers"], proveedor, hora="10:30")
    assert repetida.status_code == 409
    assert "ya fue apartado" in repetida.json()["detail"]


def test_agendar_con_proveedor_inexistente_regresa_404(client, proveedor):
    r = _agendar(client, proveedor["headers"], {"clave": "PRV-999999"})
    assert r.status_code == 404


def test_no_se_agenda_en_el_pasado(client, proveedor):
    ayer = (date.today() - timedelta(days=1)).isoformat()
    assert _agendar(client, proveedor["headers"], proveedor, fecha=ayer).status_code == 400


def test_no_se_agenda_mas_alla_de_la_ventana_permitida(client, proveedor):
    lejos = (date.today() + timedelta(days=DIAS_CITA_MAXIMO + 1)).isoformat()
    assert _agendar(client, proveedor["headers"], proveedor, fecha=lejos).status_code == 400


def test_se_puede_agendar_a_cualquier_hora_si_el_medico_no_tiene_horario(client, proveedor):
    """El médico puede aceptar por teléfono las 12:45 aunque la lista sugiera 12:00 y 13:30."""
    manana = (date.today() + timedelta(days=1)).isoformat()   # mañana cualquier hora vale
    for hora in ("12:45", "00:00", "03:15", "23:59"):
        r = _agendar(client, proveedor["headers"], proveedor, hora=hora, fecha=manana)
        assert r.status_code == 200, (hora, r.text)
    assert {c["hora"] for c in _citas(client, proveedor["headers"])} == {"12:45", "00:00", "03:15", "23:59"}


def test_rechaza_una_hora_mal_escrita(client, proveedor):
    for hora in ("25:00", "12:60", "ab:cd", "12-45"):
        r = _agendar(client, proveedor["headers"], proveedor, hora=hora)
        assert r.status_code == 400, hora
        assert "HH:MM" in r.json()["detail"]
    # Con menos de 5 caracteres ni siquiera llega a esa regla: lo frena el formato.
    assert _agendar(client, proveedor["headers"], proveedor, hora="9:30").status_code == 422


def test_se_puede_agendar_en_una_fecha_lejana_dentro_del_limite(client, proveedor):
    lejos = (date.today() + timedelta(days=100)).isoformat()   # antes el tope era de 30 días
    assert _agendar(client, proveedor["headers"], proveedor, hora="12:45", fecha=lejos).status_code == 200


def test_una_hora_de_hoy_que_ya_paso_se_rechaza_pero_mañana_a_esa_hora_si(client, proveedor, monkeypatch):
    monkeypatch.setattr(
        agenda_proveedores_service, "ahora", lambda: datetime.combine(date.today(), time(15, 0))
    )
    ya_paso = _agendar(client, proveedor["headers"], proveedor, hora="09:00")
    assert ya_paso.status_code == 400 and "ya pasó" in ya_paso.json()["detail"]
    assert _agendar(client, proveedor["headers"], proveedor, hora="15:00").status_code == 400   # ya empezó
    assert _agendar(client, proveedor["headers"], proveedor, hora="15:01").status_code == 200
    manana = (date.today() + timedelta(days=1)).isoformat()
    assert _agendar(client, proveedor["headers"], proveedor, hora="09:00", fecha=manana).status_code == 200


def test_dos_citas_en_el_mismo_minuto_con_el_mismo_medico_se_rechazan_pero_a_otro_minuto_no(client, proveedor):
    assert _agendar(client, proveedor["headers"], proveedor, hora="12:45").status_code == 200
    assert _agendar(client, proveedor["headers"], proveedor, hora="12:45").status_code == 409
    assert _agendar(client, proveedor["headers"], proveedor, hora="12:50").status_code == 200


# ---------- Agenda (disponibilidad real) ----------

def test_la_agenda_deja_de_ofrecer_el_horario_ya_apartado(client, proveedor):
    url = f"/directorio-proveedores/{proveedor['clave']}/agenda"
    agenda = client.get(url, headers=proveedor["headers"]).json()
    assert agenda[0]["fecha"] == HOY
    assert "09:00" in agenda[0]["horarios"]

    _agendar(client, proveedor["headers"], proveedor, hora="09:00")

    agenda = client.get(url, headers=proveedor["headers"]).json()
    assert "09:00" not in agenda[0]["horarios"]
    assert len(agenda[0]["horarios"]) == len(HORARIOS_BASE) - 1
    assert agenda[0]["ocupados"] == 1


def test_la_agenda_de_un_proveedor_no_afecta_la_de_otro(client, proveedor):
    """Cada renglón del padrón tiene su propia agenda -- incluido el mismo
    médico con otra especialidad."""
    headers = proveedor["headers"]
    padron = _cargar_proveedores(client, headers, [{"nombre": "Dr. Segundo", "especialidad": "Pediatría", "cedula": "1112223"}])
    otra_clave = next(p["clave"] for p in padron if p["clave"] != proveedor["clave"])

    _agendar(client, headers, proveedor, hora="09:00")

    agenda_otro = client.get(f"/directorio-proveedores/{otra_clave}/agenda", headers=headers).json()
    assert "09:00" in agenda_otro[0]["horarios"]


def test_agenda_de_proveedor_inexistente_regresa_404(client, proveedor):
    r = client.get("/directorio-proveedores/PRV-999999/agenda", headers=proveedor["headers"])
    assert r.status_code == 404


def test_la_agenda_de_un_dia_lejano_se_consulta_con_desde(client, proveedor):
    dia = (date.today() + timedelta(days=90)).isoformat()
    r = client.get(
        f"/directorio-proveedores/{proveedor['clave']}/agenda", params={"desde": dia, "dias": 1},
        headers=proveedor["headers"],
    )
    assert r.status_code == 200
    cuerpo = r.json()
    assert len(cuerpo) == 1 and cuerpo[0]["fecha"] == dia
    assert cuerpo[0]["horarios"] == list(HORARIOS_BASE)
    assert cuerpo[0]["atiende"] is True and cuerpo[0]["segun_horario"] is False
    fuera = (date.today() + timedelta(days=DIAS_CITA_MAXIMO + 1)).isoformat()
    assert client.get(
        f"/directorio-proveedores/{proveedor['clave']}/agenda", params={"desde": fuera}, headers=proveedor["headers"]
    ).status_code == 400


def test_la_agenda_de_hoy_ya_no_ofrece_los_horarios_que_pasaron(client, proveedor, monkeypatch):
    monkeypatch.setattr(
        agenda_proveedores_service, "ahora", lambda: datetime.combine(date.today(), time(12, 15))
    )
    hoy = client.get(f"/directorio-proveedores/{proveedor['clave']}/agenda", headers=proveedor["headers"]).json()[0]
    assert hoy["fecha"] == HOY
    assert hoy["horarios"] == ["13:30", "16:00", "17:30", "19:00"]


# ---------- Horario de atención del médico ----------
# Un médico que consulta de 15:00 a 18:00 no recibe sugerencias ni citas a las 09:00.

def _proximo(dia_semana: int) -> str:
    """La próxima fecha (a partir de mañana) que cae en ese día (0 = lunes)."""
    dia = date.today() + timedelta(days=1)
    while dia.weekday() != dia_semana:
        dia += timedelta(days=1)
    return dia.isoformat()


@pytest.fixture()
def con_horario(client, proveedor):
    padron = _cargar_proveedores(client, proveedor["headers"], [{
        "nombre": "Dr. Tardes", "especialidad": "Pediatría", "cedula": "7777777",
        "horario": "Lun-Vie 15:00-18:00; Sáb 09:00-11:00",
    }])
    clave = next(p["clave"] for p in padron if p["nombre"] == "Dr. Tardes")
    return {"headers": proveedor["headers"], "clave": clave}


def test_los_horarios_sugeridos_salen_del_horario_del_medico(client, con_horario):
    lunes, sabado, domingo = _proximo(0), _proximo(5), _proximo(6)
    url = f"/directorio-proveedores/{con_horario['clave']}/agenda"
    dia_lunes = client.get(url, params={"desde": lunes, "dias": 1}, headers=con_horario["headers"]).json()[0]
    assert dia_lunes["horarios"] == ["15:00", "15:30", "16:00", "16:30", "17:00", "17:30"]
    assert dia_lunes["atiende"] is True and dia_lunes["segun_horario"] is True
    dia_sabado = client.get(url, params={"desde": sabado, "dias": 1}, headers=con_horario["headers"]).json()[0]
    assert dia_sabado["horarios"] == ["09:00", "09:30", "10:00", "10:30"]
    dia_domingo = client.get(url, params={"desde": domingo, "dias": 1}, headers=con_horario["headers"]).json()[0]
    assert dia_domingo["horarios"] == [] and dia_domingo["atiende"] is False


def test_no_se_agenda_fuera_del_horario_del_medico(client, con_horario):
    lunes, domingo = _proximo(0), _proximo(6)
    temprano = _agendar(client, con_horario["headers"], con_horario, hora="09:00", fecha=lunes)
    assert temprano.status_code == 400
    assert "fuera del horario" in temprano.json()["detail"] and "15:00-18:00" in temprano.json()["detail"]
    justo_al_cerrar = _agendar(client, con_horario["headers"], con_horario, hora="18:00", fecha=lunes)
    assert justo_al_cerrar.status_code == 400
    no_atiende = _agendar(client, con_horario["headers"], con_horario, hora="16:00", fecha=domingo)
    assert no_atiende.status_code == 400 and "no atiende ese día" in no_atiende.json()["detail"]
    # Dentro del horario, a cualquier minuto (no solo en los sugeridos).
    assert _agendar(client, con_horario["headers"], con_horario, hora="16:45", fecha=lunes).status_code == 200


def test_un_horario_que_no_se_puede_leer_no_bloquea_nada(client, proveedor):
    padron = _cargar_proveedores(client, proveedor["headers"], [{
        "nombre": "Dr. Texto Libre", "especialidad": "Pediatría", "cedula": "8888888",
        "horario": "de lunes a viernes por la tarde",
    }])
    clave = next(p["clave"] for p in padron if p["nombre"] == "Dr. Texto Libre")
    lunes = _proximo(0)
    dia = client.get(f"/directorio-proveedores/{clave}/agenda", params={"desde": lunes, "dias": 1},
                     headers=proveedor["headers"]).json()[0]
    assert dia["horarios"] == list(HORARIOS_BASE) and dia["segun_horario"] is False
    assert _agendar(client, proveedor["headers"], {"clave": clave}, hora="09:00", fecha=lunes).status_code == 200


@pytest.mark.parametrize("horario, esperado", [
    ("Lun-Vie 09:00-14:00", {d: [("09:00", "14:00")] for d in range(5)}),
    ("Lun, Mié, Vie 16:00-19:30", {0: [("16:00", "19:30")], 2: [("16:00", "19:30")], 4: [("16:00", "19:30")]}),
    ("Lun-Mar 09:00-13:00; Lun 16:00-18:00", {0: [("09:00", "13:00"), ("16:00", "18:00")], 1: [("09:00", "13:00")]}),
    ("sab 09:00 - 13:00", {5: [("09:00", "13:00")]}),
    ("", None), (None, None), ("Lunes a viernes", None), ("Vie-Lun 09:00-14:00", None), ("Lun 14:00-09:00", None),
])
def test_lectura_del_horario(horario, esperado):
    assert agenda_proveedores_service.horario_por_dia(horario) == esperado


# ---------- Cancelar ----------

def test_cancelar_libera_el_horario(client, proveedor):
    folio = _agendar(client, proveedor["headers"], proveedor, hora="16:00").json()["folio"]
    assert client.post(f"/directorio-proveedores/citas/{folio}/cancelar", headers=proveedor["headers"]).status_code == 200
    agenda = client.get(f"/directorio-proveedores/{proveedor['clave']}/agenda", headers=proveedor["headers"]).json()
    assert "16:00" in agenda[0]["horarios"]


def test_cancelar_una_cita_inexistente_regresa_404(client, proveedor):
    r = client.post("/directorio-proveedores/citas/CITA-99999/cancelar", headers=proveedor["headers"])
    assert r.status_code == 404


def test_una_cita_vieja_sin_expediente_solo_la_cancela_el_administrador(client, proveedor, prov_headers):
    """Antes de que el expediente fuera obligatorio, cualquiera con el folio
    podía cancelar una cita suelta."""
    from app.repositories import citas_proveedor_repo

    vieja = citas_proveedor_repo.agendar(proveedor["clave"], HOY, "11:00", "Ana", "RFCADMTEST", None)
    r = client.post(f"/directorio-proveedores/citas/{vieja['folio']}/cancelar", headers=prov_headers)
    assert r.status_code == 403
    assert client.post(
        f"/directorio-proveedores/citas/{vieja['folio']}/cancelar", headers=proveedor["headers"]
    ).status_code == 200


# ---------- Permisos ----------

def test_proveedor_y_administrador_agendan_cabina_no(client, proveedor, prov_headers):
    assert _agendar(client, prov_headers, proveedor, hora="09:00").status_code == 200
    assert _agendar(client, proveedor["headers"], proveedor, hora="10:30").status_code == 200
    assert _agendar(client, _auth_headers_cabina(client), proveedor, hora="12:00").status_code == 403


# ---------- La cita y su expediente ----------

def test_la_cita_toma_el_paciente_del_expediente_y_devuelve_el_texto(client, proveedor, prov_headers):
    r = client.post("/directorio-proveedores/citas", headers=prov_headers, json={
        "clave_proveedor": proveedor["clave"], "fecha": HOY, "hora": "10:30", "cl_expediente": 5001,
        "paciente": "Nombre que mandó el navegador",
    })
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["cl_expediente"] == 5001
    texto = cuerpo["observaciones_sugeridas"]
    # El paciente sale del expediente, no de lo que manda el navegador.
    assert "Paciente: Juan de Dios Peralta" in texto
    assert "Nombre que mandó el navegador" not in texto
    # Trae el médico, su especialidad, el contacto, la fecha y hora y el folio.
    assert texto.startswith(f"CITA COORDINADA - Folio {cuerpo['folio']}")
    assert "Médico: Dra. Laura Méndez Ruiz" in texto
    assert "Especialidad: Cardiología" in texto
    assert "Teléfono: No registrado" in texto and "Correo: No registrado" in texto
    assert "10:30 hrs" in texto
    assert proveedor["clave"] in texto
    # Lo que no se capturó no se imprime como renglón vacío.
    assert "Costo de consulta" not in texto and "Convenio:" not in texto and "Horario de atención" not in texto
    # No se agrega nada sobre el tipo de proveedor (decisión del usuario, 2026-10-26).
    assert "Alta rápida" not in texto and "red" not in texto.lower()

    guardada = _citas(client, prov_headers)[0]
    assert guardada["cl_expediente"] == 5001 and guardada["paciente"] == "Juan de Dios Peralta"


def test_un_expediente_solo_puede_tener_una_cita_activa(client, proveedor, prov_headers):
    assert _agendar(client, prov_headers, proveedor, hora="09:00", expediente=5001).status_code == 200
    segunda = _agendar(client, prov_headers, proveedor, hora="12:00", expediente=5001)
    assert segunda.status_code == 409
    assert "ya tiene una cita agendada" in segunda.json()["detail"]
    assert "Dra. Laura Méndez Ruiz" in segunda.json()["detail"]
    # La segunda no dejó nada guardado ni apartó el horario.
    assert len(_citas(client, prov_headers)) == 1
    agenda = client.get(f"/directorio-proveedores/{proveedor['clave']}/agenda", headers=prov_headers).json()
    assert "12:00" in agenda[0]["horarios"]


def test_para_cambiar_la_cita_primero_se_cancela_y_luego_se_agenda_otra(client, proveedor, prov_headers):
    primera = _agendar(client, prov_headers, proveedor, hora="09:00", expediente=5001).json()
    assert client.post(f"/directorio-proveedores/citas/{primera['folio']}/cancelar", headers=prov_headers).status_code == 200
    assert client.get("/directorio-proveedores/citas/expediente/5001", headers=prov_headers).json()["cita"] is None

    nueva = _agendar(client, prov_headers, proveedor, hora="12:00", expediente=5001)
    assert nueva.status_code == 200
    assert nueva.json()["folio"] != primera["folio"]
    vigente = client.get("/directorio-proveedores/citas/expediente/5001", headers=prov_headers).json()["cita"]
    assert vigente["folio"] == nueva.json()["folio"]
    assert vigente["hora"] == "12:00"
    assert "Folio " + nueva.json()["folio"] in vigente["observaciones_sugeridas"]


def test_el_expediente_tiene_que_seguir_en_una_etapa_donde_se_actua(client, proveedor, prov_headers, monkeypatch):
    for estatus in ("Seguimiento Proveedor", "Finalizado", "Abierto"):
        _falso_expediente(monkeypatch, estatus=estatus)
        assert _agendar(client, prov_headers, proveedor).status_code == 400, estatus
    # También en "Seguimiento de Cita" (segunda vuelta) se puede coordinar.
    _falso_expediente(monkeypatch, estatus="Seguimiento de Cita")
    assert _agendar(client, prov_headers, proveedor).status_code == 200


def test_con_un_expediente_inexistente_regresa_404(client, proveedor, prov_headers, monkeypatch):
    _falso_expediente(monkeypatch, existe=False)
    assert _agendar(client, prov_headers, proveedor).status_code == 404


def test_no_se_cancela_una_cita_de_expediente_que_ya_salio_de_sus_manos(client, proveedor, prov_headers, monkeypatch):
    cita = _agendar(client, prov_headers, proveedor).json()
    # El Proveedor ya mandó su seguimiento: el expediente pasó a revisión del coordinador.
    _falso_expediente(monkeypatch, estatus="Seguimiento Proveedor")
    r = client.post(f"/directorio-proveedores/citas/{cita['folio']}/cancelar", headers=prov_headers)
    assert r.status_code == 400
    assert _citas(client, prov_headers)[0]["estatus"] == "Agendada"


def test_consultar_la_cita_de_un_expediente_no_filtra_el_paciente_a_quien_no_tiene_acceso(client, proveedor, prov_headers, monkeypatch):
    _agendar(client, prov_headers, proveedor, expediente=5001)
    _falso_expediente(monkeypatch, estatus="Finalizado")   # ya no es de su incumbencia
    r = client.get("/directorio-proveedores/citas/expediente/5001", headers=prov_headers)
    assert r.status_code == 400 and "Juan" not in r.text


def test_el_texto_de_la_cita_es_lo_mas_completo_posible(client):
    from app.services.cita_texto_service import observaciones_de_cita

    proveedor = {
        "clave": "PRV-000007", "nombre": "Dr. Completo", "especialidad": "CARDIOLOGIA", "consultorio": "Consultorio 302",
        "calle_num_ext_int": "Av. Vallarta 1450, Int. 302", "colonia": "Americana", "codigo_postal": "6700",
        "municipio": "Guadalajara", "estado": "Jalisco", "telefono": "33 1234 5678", "email": "dr@correo.mx",
        "horario": "Lun-Vie 09:00-14:00", "costo_consulta": "800", "convenio": "GNP", "teleconsulta": 1,
    }
    cita = {"folio": "CITA-00009", "fecha": "2026-10-26", "hora": "10:30"}
    texto = observaciones_de_cita(proveedor, cita, "Juan Pérez")
    for esperado in (
        "Folio CITA-00009", "Paciente: Juan Pérez", "Médico: Dr. Completo", "Especialidad: CARDIOLOGIA",
        "Consultorio: Consultorio 302", "Domicilio: Av. Vallarta 1450, Int. 302, Americana, C.P. 06700, Guadalajara, Jalisco",
        "Teléfono: 33 1234 5678", "Correo: dr@correo.mx", "Fecha y hora de la cita: Lun 26/10/2026, 10:30 hrs",
        "Horario de atención: Lun-Vie 09:00-14:00", "Costo de consulta: $800", "Convenio: GNP", "Teleconsulta: Sí",
        "Clave del proveedor: PRV-000007",
    ):
        assert esperado in texto, esperado
    assert len(texto) < 1000


def test_el_costo_por_rango_se_imprime_como_aproximado():
    from app.services.cita_texto_service import observaciones_de_cita

    texto = observaciones_de_cita({"costo_consulta": "600-900"}, {"folio": "CITA-1", "fecha": "2026-10-26", "hora": "10:30"}, None)
    assert "Costo de consulta: $600 a $900 (aproximado)" in texto


# ---------- Elegir expediente al agendar desde el directorio ----------

def test_expedientes_para_cita_solo_trae_los_accionables_y_sin_cita(client, proveedor, prov_headers, monkeypatch):
    """Lista del paso "¿Para qué expediente?". El estatus vive en SQLite; los
    datos del expediente, en el sistema central (aquí simulado): los recientes
    salen de una sola consulta y los más viejos se piden uno por uno."""
    import app.repositories.expedientes_repo as expedientes_repo
    from app.repositories import estatus_repo
    from app.schemas.expediente import EstatusExpediente

    estatus_repo.actualizar_estatus(5001, EstatusExpediente.EN_ESPERA_RESPUESTA, "RFCCAB")
    estatus_repo.actualizar_estatus(5002, EstatusExpediente.SEGUIMIENTO_CITA, "RFCCAB")   # tendrá cita
    estatus_repo.actualizar_estatus(5003, EstatusExpediente.SEGUIMIENTO_COORDINADOR, "RFCCAB")   # Finalizado
    estatus_repo.actualizar_estatus(5004, EstatusExpediente.EN_ESPERA_RESPUESTA, "RFCCAB")   # viejo
    _agendar(client, prov_headers, proveedor, expediente=5002)

    consultas = []

    def listar(filtro):
        consultas.append(filtro.cl_expediente)
        def con_su_estatus(cl):
            detalle = estatus_repo.obtener_detalle(cl)
            return _expediente(cl, EstatusExpediente(detalle["estatus"]) if detalle else EstatusExpediente.ABIERTO)
        if filtro.cl_expediente:
            return [con_su_estatus(filtro.cl_expediente)]
        return [con_su_estatus(cl) for cl in (5001, 5002, 5003)]   # el 5004 ya no sale en el rango reciente

    monkeypatch.setattr(expedientes_repo, "listar_expedientes", listar)
    r = client.get("/directorio-proveedores/expedientes-para-cita", headers=prov_headers)
    assert r.status_code == 200, r.text
    assert [e["expediente"] for e in r.json()] == [5004, 5001]
    assert r.json()[0]["nombre_paciente"] == "Juan de Dios Peralta"
    assert consultas == [None, 5004], "una consulta del rango y solo el viejo uno por uno"

    assert client.get("/directorio-proveedores/expedientes-para-cita", headers=proveedor["headers"]).status_code == 200
    assert client.get("/directorio-proveedores/expedientes-para-cita",
                      headers=_auth_headers_cabina(client)).status_code == 403
