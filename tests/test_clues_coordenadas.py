"""
Pruebas del enriquecimiento de coordenadas con el Catálogo CLUES
(app/services/clues_service.py y proveedores_repo.asignar_coordenadas).

No se descarga el archivo real de 25 MB: se arma uno pequeño con el mismo
layout. Lo que importa comprobar es que NO se asignen coordenadas equivocadas
-- el RFC de CLUES es de la persona moral y una cadena tiene el mismo RFC en
decenas de sucursales.
"""
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.services import clues_service
from app.services.clues_service import ErrorCatalogoClues

# Encabezado real del archivo de la DGIS, recortado a las columnas que se leen.
ENCABEZADO = [
    "CLUES", "NOMBRE DE LA INSTITUCION", "ESTATUS DE OPERACION",
    "NOMBRE DE LA UNIDAD", "NOMBRE COMERCIAL", "RFC DEL ESTABLECIMIENTO",
    "ENTIDAD", "MUNICIPIO", "CODIGO POSTAL", "NOMBRE DE TIPOLOGIA",
    "LATITUD", "LONGITUD",
]

PRIVADA = "SERVICIOS MEDICOS PRIVADOS"
OPERANDO = "EN OPERACION"


def _fila(**kw):
    base = {
        "CLUES": "JCSSA000001", "NOMBRE DE LA INSTITUCION": PRIVADA,
        "ESTATUS DE OPERACION": OPERANDO, "NOMBRE DE LA UNIDAD": "CENTRO MEDICO",
        "NOMBRE COMERCIAL": "CENTRO MEDICO", "RFC DEL ESTABLECIMIENTO": "CME-010203-AB1",
        "ENTIDAD": "JALISCO", "MUNICIPIO": "GUADALAJARA", "CODIGO POSTAL": 44100,
        "NOMBRE DE TIPOLOGIA": "CONSULTORIO PARTICULAR",
        "LATITUD": 20.6767, "LONGITUD": -103.3475,
    }
    base.update(kw)
    return [base[c] for c in ENCABEZADO]


def _archivo(filas, hoja="CLUES_202608"):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(ENCABEZADO)
    for f in filas:
        ws.append(f)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _proveedor(**kw):
    base = {
        "clave": "PRV-000001", "nombre": "CENTRO MEDICO", "rfc": "CME010203AB1",
        "estado": "Jalisco", "municipio": "Guadalajara", "codigo_postal": "44100",
        "latitud": None, "longitud": None,
    }
    base.update(kw)
    return base


# ---------- Lectura del catálogo ----------

def test_lee_los_establecimientos_privados_en_operacion():
    archivo = _archivo([
        _fila(),
        _fila(CLUES="JCSSA000002", **{"NOMBRE DE LA INSTITUCION": "SECRETARIA DE SALUD"}),
        _fila(CLUES="JCSSA000003", **{"ESTATUS DE OPERACION": "FUERA DE OPERACION"}),
    ])
    establecimientos = clues_service.leer_establecimientos(archivo)
    assert len(establecimientos) == 1, "solo el privado en operación"
    assert establecimientos[0]["clues"] == "JCSSA000001"


def test_la_hoja_se_localiza_por_prefijo():
    """El nombre lleva el corte ("CLUES_202608") y cambia en cada publicación."""
    establecimientos = clues_service.leer_establecimientos(
        _archivo([_fila()], hoja="CLUES_210112")
    )
    assert len(establecimientos) == 1


def test_sin_la_hoja_de_establecimientos_se_avisa():
    wb = Workbook()
    wb.active.title = "OtraCosa"
    wb.active.append(ENCABEZADO)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    with pytest.raises(ErrorCatalogoClues, match="hoja de establecimientos"):
        clues_service.leer_establecimientos(buffer)


def test_si_faltan_columnas_se_avisa():
    wb = Workbook()
    wb.active.title = "CLUES_202608"
    wb.active.append(["CLUES", "MUNICIPIO"])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    with pytest.raises(ErrorCatalogoClues, match="faltan columnas"):
        clues_service.leer_establecimientos(buffer)


def test_el_codigo_postal_se_normaliza_a_cinco_digitos():
    """Excel lo guarda como número y le come el cero de la izquierda."""
    establecimientos = clues_service.leer_establecimientos(
        _archivo([_fila(**{"CODIGO POSTAL": 6700})])
    )
    assert establecimientos[0]["codigo_postal"] == "06700"


# ---------- El cruce ----------

def test_empata_por_rfc_aunque_este_escrito_distinto():
    """CLUES trae "CME-010203-AB1" y el padrón "CME010203AB1"."""
    establecimientos = clues_service.leer_establecimientos(_archivo([_fila()]))
    plan = clues_service.emparejar([_proveedor()], establecimientos)
    assert len(plan["asignaciones"]) == 1
    asignado = plan["asignaciones"][0]
    assert asignado["cruce"] == clues_service.CRUCE_RFC
    assert asignado["latitud"] == pytest.approx(20.6767)


def test_empata_por_nombre_cuando_el_rfc_no_sirve():
    establecimientos = clues_service.leer_establecimientos(_archivo([_fila()]))
    plan = clues_service.emparejar([_proveedor(rfc=None)], establecimientos)
    assert len(plan["asignaciones"]) == 1
    assert plan["asignaciones"][0]["cruce"] == clues_service.CRUCE_NOMBRE


def test_los_acentos_y_mayusculas_no_impiden_el_cruce():
    establecimientos = clues_service.leer_establecimientos(
        _archivo([_fila(MUNICIPIO="MERIDA", ENTIDAD="YUCATAN")])
    )
    plan = clues_service.emparejar(
        [_proveedor(municipio="Mérida", estado="Yucatán")], establecimientos
    )
    assert len(plan["asignaciones"]) == 1


# ---------- Lo que NO debe pasar ----------

def test_no_asigna_coordenadas_de_otra_sucursal_con_el_mismo_rfc():
    """
    EL CASO QUE MOTIVÓ EL DISEÑO: el RFC de CLUES es de la persona moral. Con
    el archivo real, cruzar solo por RFC le ponía a 12 proveedores las
    coordenadas de un establecimiento de OTRO municipio.
    """
    establecimientos = clues_service.leer_establecimientos(_archivo([
        _fila(CLUES="A1", MUNICIPIO="MONTERREY", **{"CODIGO POSTAL": 64000},
              LATITUD=25.6866, LONGITUD=-100.3161),
    ]))
    # Mismo RFC, pero el proveedor está en Guadalajara.
    plan = clues_service.emparejar([_proveedor(municipio="Guadalajara")], establecimientos)
    assert plan["asignaciones"] == [], "no debe tomar las coordenadas de Monterrey"
    assert len(plan["sin_coincidencia"]) == 1


def test_un_rfc_repetido_en_la_misma_zona_queda_como_ambiguo():
    """Dos sucursales en el mismo municipio y CP: no hay forma de saber cuál
    es, así que no se asigna ninguna."""
    establecimientos = clues_service.leer_establecimientos(_archivo([
        _fila(CLUES="A1", LATITUD=20.60, LONGITUD=-103.30),
        _fila(CLUES="A2", LATITUD=20.70, LONGITUD=-103.40),
    ]))
    plan = clues_service.emparejar([_proveedor()], establecimientos)
    assert plan["asignaciones"] == []


def test_no_propone_nada_para_quien_ya_tiene_coordenadas():
    """Las del Excel del área son la fuente autorizada."""
    establecimientos = clues_service.leer_establecimientos(_archivo([_fila()]))
    plan = clues_service.emparejar(
        [_proveedor(latitud=19.0, longitud=-99.0)], establecimientos
    )
    assert plan["asignaciones"] == []
    assert plan["ya_tenian_coordenadas"] == 1


def test_ignora_los_establecimientos_sin_coordenadas():
    establecimientos = clues_service.leer_establecimientos(
        _archivo([_fila(LATITUD=None, LONGITUD=None)])
    )
    plan = clues_service.emparejar([_proveedor()], establecimientos)
    assert plan["asignaciones"] == []
    assert plan["establecimientos_con_coordenadas"] == 0


def test_un_rfc_incompleto_no_identifica():
    """Menos de 12 caracteres es un RFC a medias; si se aceptara, cruzaría con
    cualquiera que empiece igual."""
    establecimientos = clues_service.leer_establecimientos(
        _archivo([_fila(**{"RFC DEL ESTABLECIMIENTO": "CME-01"})])
    )
    plan = clues_service.emparejar(
        [_proveedor(rfc="CME01", nombre="OTRO NOMBRE")], establecimientos
    )
    assert plan["asignaciones"] == []


# ---------- Escritura en el padrón ----------

@pytest.fixture()
def base(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_AUTH_MODE", "sql")
    monkeypatch.setenv("DB_SERVER", "fake")
    monkeypatch.setenv("DB_NAME", "fake")
    monkeypatch.setenv("DB_USER", "fake")
    monkeypatch.setenv("DB_PASSWORD", "fake")
    monkeypatch.setenv("LOCAL_DB_PATH", str(tmp_path / "clues.db"))
    monkeypatch.setenv("JWT_SECRET_KEY", "clave-de-prueba")

    from app.config import get_settings

    get_settings.cache_clear()
    from app.db.local_store import init_local_db

    init_local_db()
    yield
    get_settings.cache_clear()


def _cargar(registros):
    from app.repositories import proveedores_repo

    return proveedores_repo.cargar_padron(registros, "ADMI000001")


def _registro(**kw):
    base_reg = {
        "nombre": "CENTRO MEDICO", "especialidad": "CARDIOLOGIA", "rfc": "CME010203AB1",
        "estado": "Jalisco", "municipio": "Guadalajara", "codigo_postal": "44100",
        "calle_num_ext_int": "Av. Vallarta 100", "telefono": "3312345678",
        "latitud": None, "longitud": None,
    }
    base_reg.update(kw)
    return base_reg


def test_asignar_coordenadas_las_guarda_y_marca_su_origen(base):
    from app.repositories import proveedores_repo

    _cargar([_registro()])
    clave = proveedores_repo.listar_padron()[0]["clave"]

    cambiados = proveedores_repo.asignar_coordenadas(
        [{"clave": clave, "latitud": 20.6767, "longitud": -103.3475}]
    )
    assert cambiados == 1

    proveedor = proveedores_repo.obtener_por_clave(clave)
    assert proveedor["latitud"] == pytest.approx(20.6767)
    assert proveedor["origen_coordenadas"] == proveedores_repo.ORIGEN_COORDENADAS_CLUES


def test_asignar_coordenadas_no_pisa_las_que_vinieron_en_la_carga(base):
    """Si el Excel del área ya las traía, mandan esas."""
    from app.repositories import proveedores_repo

    _cargar([_registro(latitud=19.4326, longitud=-99.1332)])
    clave = proveedores_repo.listar_padron()[0]["clave"]
    assert proveedores_repo.obtener_por_clave(clave)["origen_coordenadas"] == \
        proveedores_repo.ORIGEN_COORDENADAS_CARGA

    cambiados = proveedores_repo.asignar_coordenadas(
        [{"clave": clave, "latitud": 20.6767, "longitud": -103.3475}]
    )
    assert cambiados == 0, "no debió tocar al que ya tenía coordenadas"
    assert proveedores_repo.obtener_por_clave(clave)["latitud"] == pytest.approx(19.4326)


def test_una_carga_posterior_sin_coordenadas_no_borra_el_enriquecimiento(base):
    """
    El Excel del área entrega las coordenadas VACÍAS. Como la carga es
    incremental y refresca los datos, sin este cuidado la siguiente carga
    dejaba en NULL lo que CLUES ya había llenado.
    """
    from app.repositories import proveedores_repo

    _cargar([_registro()])
    clave = proveedores_repo.listar_padron()[0]["clave"]
    proveedores_repo.asignar_coordenadas(
        [{"clave": clave, "latitud": 20.6767, "longitud": -103.3475}]
    )

    # Vuelve a llegar el MISMO proveedor en una carga, otra vez sin coordenadas.
    resultado = _cargar([_registro(costo_consulta="950")])
    assert resultado["actualizados"] == 1

    proveedor = proveedores_repo.obtener_por_clave(clave)
    assert proveedor["latitud"] == pytest.approx(20.6767), "se conservan las de CLUES"
    assert proveedor["costo_consulta"] == "950", "y sí se actualizó lo demás"


def test_una_carga_con_coordenadas_si_gana_sobre_las_de_clues(base):
    from app.repositories import proveedores_repo

    _cargar([_registro()])
    clave = proveedores_repo.listar_padron()[0]["clave"]
    proveedores_repo.asignar_coordenadas(
        [{"clave": clave, "latitud": 20.6767, "longitud": -103.3475}]
    )

    _cargar([_registro(latitud=19.4326, longitud=-99.1332)])
    proveedor = proveedores_repo.obtener_por_clave(clave)
    assert proveedor["latitud"] == pytest.approx(19.4326)
    assert proveedor["origen_coordenadas"] == proveedores_repo.ORIGEN_COORDENADAS_CARGA


def test_contar_coordenadas_resume_el_estado_del_padron(base):
    from app.repositories import proveedores_repo

    _cargar([
        _registro(nombre="CON COORDENADAS", latitud=19.4, longitud=-99.1, telefono="1"),
        _registro(nombre="SIN COORDENADAS", telefono="2"),
        _registro(nombre="PARA ENRIQUECER", telefono="3"),
    ])
    clave = next(
        p["clave"] for p in proveedores_repo.listar_padron() if p["nombre"] == "PARA ENRIQUECER"
    )
    proveedores_repo.asignar_coordenadas([{"clave": clave, "latitud": 20.6, "longitud": -103.3}])

    resumen = proveedores_repo.contar_coordenadas()
    assert resumen["total"] == 3
    assert resumen["con_coordenadas"] == 2
    assert resumen["sin_coordenadas"] == 1
    assert resumen["por_origen"] == {"carga": 1, "clues": 1}


def test_el_flujo_completo_enriquece_el_padron(base):
    """De punta a punta: catálogo -> cruce -> escritura."""
    from app.repositories import proveedores_repo

    _cargar([
        _registro(nombre="CENTRO MEDICO", telefono="1"),
        _registro(nombre="CLINICA SIN CLUES", rfc="XXX010203AB1", telefono="2"),
    ])
    establecimientos = clues_service.leer_establecimientos(_archivo([_fila()]))
    padron = proveedores_repo.listar_padron()

    plan = clues_service.emparejar(padron, establecimientos)
    assert len(plan["asignaciones"]) == 1
    assert len(plan["sin_coincidencia"]) == 1

    assert proveedores_repo.asignar_coordenadas(plan["asignaciones"]) == 1
    resumen = proveedores_repo.contar_coordenadas()
    assert resumen["con_coordenadas"] == 1
    assert resumen["por_origen"] == {"clues": 1}
