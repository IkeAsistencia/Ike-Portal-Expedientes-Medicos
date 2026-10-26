"""
Pruebas del catálogo de especialidades médicas y de los KPI del Directorio
de Proveedores (GET /directorio-proveedores/especialidades).

El catálogo se siembra solo al crear la base (ver
app/repositories/especialidades_repo.py), así que estas pruebas no necesitan
SQL Server ni cargar nada a mano: basta levantar el cliente.
"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.repositories import especialidades_repo
from app.services.proveedores_excel_service import COLUMNAS

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _sin_sql_server(monkeypatch):
    import app.main as main_module

    monkeypatch.setattr(main_module, "calentar_pool", lambda: None)

    from app.core.rate_limit import limpiar_todo

    limpiar_todo()


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_AUTH_MODE", "sql")
    monkeypatch.setenv("DB_SERVER", "servidor-fake")
    monkeypatch.setenv("DB_NAME", "bd-fake")
    monkeypatch.setenv("DB_USER", "usuario-fake")
    monkeypatch.setenv("DB_PASSWORD", "password-fake")
    monkeypatch.setenv("LOCAL_DB_PATH", str(tmp_path / "test_local.db"))
    monkeypatch.setenv("JWT_SECRET_KEY", "clave-de-prueba")

    from app.config import get_settings

    get_settings.cache_clear()

    with TestClient(app) as c:
        yield c


def _admin(client):
    import app.repositories.accesos_repo as a

    a.alta_acceso("RFCADMESP", "Admin", a.PERFIL_ADMINISTRADOR)
    a.crear_password("RFCADMESP", "Admin2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCADMESP", "password": "Admin2024xx"})
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def _proveedor(client):
    import app.repositories.accesos_repo as a

    a.alta_acceso("RFCPROESP", "Proveedor", a.PERFIL_PROVEEDOR, "Jalisco", "p@correo.mx")
    a.crear_password("RFCPROESP", "Proveedor2024")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCPROESP", "password": "Proveedor2024"})
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def _cabina(client):
    import app.repositories.accesos_repo as a

    a.alta_acceso("RFCCABESP", "Cabina", a.PERFIL_CABINA)
    a.crear_password("RFCCABESP", "Cabina2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCCABESP", "password": "Cabina2024xx"})
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def _xlsx(filas):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col, (encabezado, _clave, _ancho, _alias) in enumerate(COLUMNAS, start=1):
        ws.cell(row=1, column=col, value=encabezado)
    for i, fila in enumerate(filas, start=2):
        for col, (_encabezado, clave, _ancho, _alias) in enumerate(COLUMNAS, start=1):
            ws.cell(row=i, column=col, value=fila.get(clave))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _cargar(client, headers, filas):
    return client.post(
        "/proveedores-catalogo/cargar",
        headers=headers,
        files={"archivo": ("catalogo.xlsx", _xlsx(filas), XLSX_MIME)},
    )


def _medico(nombre, especialidad, **extra):
    return {
        "nombre": nombre, "especialidad": especialidad, "cedula": str(abs(hash(nombre)) % 10**7),
        "estado": "Jalisco", "municipio": "Guadalajara", **extra,
    }


# ---------- La semilla ----------

def test_el_catalogo_trae_las_38_especialidades_del_archivo(client):
    r = client.get("/directorio-proveedores/especialidades", headers=_admin(client))
    assert r.status_code == 200, r.text
    assert r.json()["total"] == len(especialidades_repo.SEMILLA) == 38


def test_la_semilla_es_idempotente(client):
    """Se resiembra en cada arranque: no debe duplicar ni perder nada."""
    headers = _admin(client)
    antes = client.get("/directorio-proveedores/especialidades", headers=headers).json()

    from app.db.local_store import init_local_db

    init_local_db()
    init_local_db()

    despues = client.get("/directorio-proveedores/especialidades", headers=headers).json()
    assert despues == antes


# ---------- El reparto entre tarjetas y combo ----------

def test_caben_16_tarjetas_en_4_renglones_de_4(client):
    cuerpo = client.get("/directorio-proveedores/especialidades", headers=_admin(client)).json()
    assert cuerpo["maximo_tarjetas"] == 16
    assert len(cuerpo["tarjetas"]) == 16, "4 renglones de 4"
    assert len(cuerpo["combo"]) == 38 - 16


def test_las_tarjetas_van_en_el_orden_del_ordenkpi(client):
    cuerpo = client.get("/directorio-proveedores/especialidades", headers=_admin(client)).json()
    ordenes = [x["orden_kpi"] for x in cuerpo["tarjetas"]]
    assert ordenes == sorted(ordenes)
    # La primera es la de ordenKPI = 1 en el archivo del usuario.
    assert cuerpo["tarjetas"][0]["descripcion"] == "MEDICINA GENERAL"
    assert cuerpo["tarjetas"][0]["orden_kpi"] == 1


def test_ninguna_especialidad_aparece_en_tarjeta_y_en_combo(client):
    cuerpo = client.get("/directorio-proveedores/especialidades", headers=_admin(client)).json()
    en_tarjeta = {x["cl_especialidad"] for x in cuerpo["tarjetas"]}
    en_combo = {x["cl_especialidad"] for x in cuerpo["combo"]}
    assert not (en_tarjeta & en_combo)
    assert len(en_tarjeta | en_combo) == cuerpo["total"], "ninguna se queda sin mostrar"


def test_las_que_no_traen_ordenkpi_van_al_final_del_combo(client):
    """CIRUGIA GENERAL y CIRUGIA ONCOLOGICA vienen con ordenKPI en NULL:
    nunca salen en tarjeta y se dejan hasta el final (decisión del usuario)."""
    cuerpo = client.get("/directorio-proveedores/especialidades", headers=_admin(client)).json()

    assert all(x["orden_kpi"] is not None for x in cuerpo["tarjetas"])

    sin_orden = [x["descripcion"] for x in cuerpo["combo"] if x["orden_kpi"] is None]
    assert sin_orden == ["CIRUGIA GENERAL", "CIRUGIA ONCOLOGICA"]
    assert [x["descripcion"] for x in cuerpo["combo"][-2:]] == sin_orden


def test_el_ordenkpi_repetido_no_cambia_de_lugar_entre_consultas(client):
    """El archivo trae el 15 dos veces (ENDOCRINOLOGIA y UROLOGIA). Sin un
    desempate fijo, podrían intercambiarse y la pantalla "bailaría"."""
    headers = _admin(client)
    primera = client.get("/directorio-proveedores/especialidades", headers=headers).json()
    segunda = client.get("/directorio-proveedores/especialidades", headers=headers).json()
    assert [x["cl_especialidad"] for x in primera["tarjetas"]] == \
           [x["cl_especialidad"] for x in segunda["tarjetas"]]

    repetidas = [x["descripcion"] for x in primera["tarjetas"] if x["orden_kpi"] == 15]
    assert repetidas == ["ENDOCRINOLOGIA", "UROLOGIA"], "desempatadas por clave"


# ---------- El KPI sale del padrón real ----------

def test_sin_padron_todas_las_especialidades_salen_en_cero(client):
    cuerpo = client.get("/directorio-proveedores/especialidades", headers=_admin(client)).json()
    for x in cuerpo["tarjetas"] + cuerpo["combo"]:
        assert x["dentro_de_red"] == 0
        assert x["fuera_de_red"] == 0
        assert x["total_proveedores"] == 0


def test_la_carga_suma_al_kpi_de_su_especialidad(client):
    headers = _admin(client)
    assert _cargar(client, headers, [
        _medico("Dra. Uno", "Cardiología"),
        _medico("Dr. Dos", "Cardiología", municipio="Zapopan"),
        _medico("Dra. Tres", "Pediatría"),
    ]).status_code == 200

    por_nombre = {
        x["descripcion"]: x
        for x in client.get("/directorio-proveedores/especialidades", headers=headers).json()["tarjetas"]
    }
    assert por_nombre["CARDIOLOGIA"]["dentro_de_red"] == 2
    assert por_nombre["CARDIOLOGIA"]["fuera_de_red"] == 0
    assert por_nombre["CARDIOLOGIA"]["zonas"] == 2, "Guadalajara y Zapopan"
    assert por_nombre["PEDIATRIA"]["dentro_de_red"] == 1
    assert por_nombre["DERMATOLOGIA"]["dentro_de_red"] == 0


def test_el_kpi_separa_dentro_y_fuera_de_la_red(client):
    """Dentro de la red = llegó en la carga; fuera = alta hecha en el portal
    y todavía sin contratar (formato que pidió el usuario)."""
    admin = _admin(client)
    _cargar(client, admin, [_medico("Dra. De la carga", "Cardiología")])

    solicitud = {
        "nombre": "Dr. Del portal", "especialidad": "CARDIOLOGIA", "cedula": "98765432",
        "rfc": "GODE561231GR8", "estado": "Jalisco", "municipio": "Guadalajara",
        "calle_num_ext_int": "Av. Vallarta 1450", "colonia": "Americana", "codigo_postal": "44100",
        "consultorio": "Consultorio del portal", "horario": "Lun-Vie 09:00-14:00",
        "telefono": "3312345678", "email": "portal@correo.mx",
    }
    folio = client.post("/solicitudes-alta-medico", json=solicitud, headers=admin).json()["folio"]
    aprobacion = client.post(
        f"/solicitudes-alta-medico/{folio}/estatus",
        json={"estatus": "Aprobada"}, headers=admin,
    )
    assert aprobacion.status_code == 200, aprobacion.text

    cardio = next(
        x for x in client.get("/directorio-proveedores/especialidades", headers=admin).json()["tarjetas"]
        if x["descripcion"] == "CARDIOLOGIA"
    )
    assert cardio["dentro_de_red"] == 1
    assert cardio["fuera_de_red"] == 1
    assert cardio["total_proveedores"] == 2


# ---------- Nombres que no empatan con el catálogo ----------

def test_el_nombre_del_padron_se_relaciona_sin_acentos_ni_mayusculas(client):
    headers = _admin(client)
    _cargar(client, headers, [
        _medico("Dra. Con acento", "Cardiología"),
        _medico("Dr. Sin acento", "CARDIOLOGIA", municipio="Zapopan"),
        _medico("Dra. Minusculas", "cardiologia", municipio="Tonalá"),
    ])
    cardio = next(
        x for x in client.get("/directorio-proveedores/especialidades", headers=headers).json()["tarjetas"]
        if x["descripcion"] == "CARDIOLOGIA"
    )
    assert cardio["dentro_de_red"] == 3


def test_los_nombres_viejos_del_portal_siguen_contando(client):
    """El padrón ya tenía "Traumatología" y "Ginecología" (los nombres cortos
    del combo fijo de antes). Sin equivalentes, su tarjeta diría 0 aunque sí
    haya a quién mostrar."""
    headers = _admin(client)
    _cargar(client, headers, [
        _medico("Dr. Trauma", "Traumatología"),
        _medico("Dra. Gine", "Ginecología"),
        _medico("Dr. Odonto", "Odontología"),
    ])
    por_nombre = {
        x["descripcion"]: x
        for x in client.get("/directorio-proveedores/especialidades", headers=headers).json()["tarjetas"]
    }
    assert por_nombre["TRAUMATOLOGIA Y ORTOPEDIA"]["dentro_de_red"] == 1
    assert por_nombre["GINECOLOGIA Y OBSTETRICIA"]["dentro_de_red"] == 1
    assert por_nombre["ODONTOLOGIA INTEGRAL"]["dentro_de_red"] == 1


def test_la_carga_reporta_las_especialidades_que_no_reconoce(client):
    """No se rechaza el archivo (decisión del usuario): el proveedor entra,
    pero hay que avisar que no va a sumar a ninguna tarjeta."""
    headers = _admin(client)
    r = _cargar(client, headers, [
        _medico("Dra. Buena", "Cardiología"),
        _medico("Dr. Raro", "Medicina Cuántica"),
        _medico("Dra. Otra", "Medicina Cuántica", municipio="Zapopan"),
    ])
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["nuevos"] == 3, "los 3 se cargan igual"
    assert cuerpo["especialidades_no_reconocidas"] == ["Medicina Cuántica"], "sin repetir"


def test_un_proveedor_con_especialidad_desconocida_no_suma_a_ninguna_tarjeta(client):
    headers = _admin(client)
    _cargar(client, headers, [_medico("Dr. Raro", "Medicina Cuántica")])

    cuerpo = client.get("/directorio-proveedores/especialidades", headers=headers).json()
    assert sum(x["total_proveedores"] for x in cuerpo["tarjetas"] + cuerpo["combo"]) == 0
    # Pero sí está en el padrón, solo que sin especialidad del catálogo.
    padron = client.get("/directorio-proveedores", headers=headers).json()
    assert len(padron) == 1
    assert padron[0]["cl_especialidad"] is None


def test_la_especialidad_queda_amarrada_en_el_padron(client):
    """proveedores.cl_especialidad es lo que hace contable el KPI."""
    headers = _admin(client)
    _cargar(client, headers, [_medico("Dra. Laura", "Cardiología")])
    padron = client.get("/directorio-proveedores", headers=headers).json()
    assert padron[0]["cl_especialidad"] == 6, "clEspecialidad de CARDIOLOGIA"


# ---------- Permisos ----------

def test_proveedor_y_administrador_ven_las_especialidades(client):
    for headers in (_admin(client), _proveedor(client)):
        r = client.get("/directorio-proveedores/especialidades", headers=headers)
        assert r.status_code == 200


def test_cabina_no_ve_las_especialidades(client):
    r = client.get("/directorio-proveedores/especialidades", headers=_cabina(client))
    assert r.status_code == 403


def test_sin_sesion_no_se_ven_las_especialidades(client):
    assert client.get("/directorio-proveedores/especialidades").status_code == 401


# ---------- Helpers del repositorio ----------

def test_normalizar_quita_acentos_y_mayusculas():
    assert especialidades_repo.normalizar("Cardiología") == "CARDIOLOGIA"
    assert especialidades_repo.normalizar("  medicina   general ") == "MEDICINA GENERAL"
    assert especialidades_repo.normalizar(None) == ""


def test_resolver_devuelve_none_si_no_existe(client):
    assert especialidades_repo.resolver("Cardiología") == 6
    assert especialidades_repo.resolver("Medicina Cuántica") is None
    assert especialidades_repo.resolver("") is None
