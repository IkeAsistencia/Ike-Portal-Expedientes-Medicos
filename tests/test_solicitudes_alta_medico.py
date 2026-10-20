"""
Pruebas de /solicitudes-alta-medico (Alta de médico, pantalla "Proveedores"
del Directorio de Proveedores) y de cómo una solicitud aprobada se vuelve
visible en GET /directorio-proveedores -- NO requiere SQL Server real: todo
vive en SQLite local, nunca se manda al sistema central.
"""
import pytest
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture(autouse=True)
def _sin_sql_server(monkeypatch):
    import app.main as main_module

    monkeypatch.setattr(main_module, "calentar_pool", lambda: None)

    from app.core.rate_limit import limpiar_todo

    limpiar_todo()


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_AUTH_MODE", "sql")
    monkeypatch.setenv("DB_SERVER", "servidor-fake")
    monkeypatch.setenv("DB_NAME", "bd-fake")
    monkeypatch.setenv("DB_USER", "usuario-fake")
    monkeypatch.setenv("DB_PASSWORD", "password-fake")
    monkeypatch.setenv("LOCAL_DB_PATH", str(tmp_path / "test_local.db"))
    monkeypatch.setenv("JWT_SECRET_KEY", "clave-de-prueba")

    from app.config import get_settings

    get_settings.cache_clear()

    with TestClient(app) as c:
        yield c


def _auth_headers_admin(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCADMTEST", "Admin de Prueba", accesos_repo_module.PERFIL_ADMINISTRADOR)
    accesos_repo_module.crear_password("RFCADMTEST", "Admin2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCADMTEST", "password": "Admin2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_proveedor(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCPROTEST", "Proveedor de Prueba", accesos_repo_module.PERFIL_PROVEEDOR)
    accesos_repo_module.crear_password("RFCPROTEST", "Proveedor2024")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCPROTEST", "password": "Proveedor2024"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_cabina(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCCABTEST", "Cabina de Prueba", accesos_repo_module.PERFIL_CABINA)
    accesos_repo_module.crear_password("RFCCABTEST", "Cabina2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCCABTEST", "password": "Cabina2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


SOLICITUD_VALIDA = {
    "nombre": "Dr. Nuevo Médico", "especialidad": "Cardiología", "cedula": "12345678",
    "rfc": "NUME900101AB1", "estado": "Jalisco", "municipio": "Guadalajara",
    "calle_num_ext_int": "Av. Siempre Viva 123, Int. 4", "colonia": "Centro",
    "codigo_postal": "44100", "telefono": "33 0000 0000", "email": "nuevo@correo.mx",
}


# ---------- Permisos (Cabina no usa esta pantalla) ----------

def test_crear_solicitud_rechaza_cabina(client):
    r = client.post("/solicitudes-alta-medico", headers=_auth_headers_cabina(client), json=SOLICITUD_VALIDA)
    assert r.status_code == 403


def test_listar_solicitudes_rechaza_cabina(client):
    r = client.get("/solicitudes-alta-medico", headers=_auth_headers_cabina(client))
    assert r.status_code == 403


def test_proveedor_y_administrador_pueden_crear_solicitud(client):
    assert client.post("/solicitudes-alta-medico", headers=_auth_headers_proveedor(client), json=SOLICITUD_VALIDA).status_code == 200
    assert client.post("/solicitudes-alta-medico", headers=_auth_headers_admin(client), json=SOLICITUD_VALIDA).status_code == 200


# ---------- Revisión: aprobar / rechazar / editar son solo de Administrador ----------
# Decisión del usuario (2026-10-19): capturar un alta lo puede hacer
# cualquier perfil menos Cabina, pero RESOLVERLA (o corregirla antes de
# resolverla) es acto de revisión y queda en Administrador -- si no, quien
# captura se aprueba a sí mismo y mete un proveedor real al padrón.


def test_proveedor_no_puede_aprobar_su_propia_solicitud(client):
    proveedor = _auth_headers_proveedor(client)
    folio = client.post("/solicitudes-alta-medico", headers=proveedor, json=SOLICITUD_VALIDA).json()["folio"]

    r = client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=proveedor, json={"estatus": "Aprobada"})
    assert r.status_code == 403
    assert "Administrador" in r.json()["detail"]

    # Y no quedó nada en el padrón.
    assert client.get("/directorio-proveedores", headers=proveedor).json() == []


def test_proveedor_no_puede_rechazar(client):
    proveedor = _auth_headers_proveedor(client)
    folio = client.post("/solicitudes-alta-medico", headers=proveedor, json=SOLICITUD_VALIDA).json()["folio"]

    r = client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=proveedor, json={"estatus": "Rechazada"})
    assert r.status_code == 403
    assert client.get("/solicitudes-alta-medico", headers=proveedor).json()[0]["estatus"] == "En revisión"


def test_proveedor_no_puede_editar_una_solicitud(client):
    proveedor = _auth_headers_proveedor(client)
    folio = client.post("/solicitudes-alta-medico", headers=proveedor, json=SOLICITUD_VALIDA).json()["folio"]

    r = client.put(
        f"/solicitudes-alta-medico/{folio}",
        headers=proveedor,
        json={**SOLICITUD_VALIDA, "nombre": "Dr. Nombre Cambiado"},
    )
    assert r.status_code == 403
    assert client.get("/solicitudes-alta-medico", headers=proveedor).json()[0]["nombre"] == SOLICITUD_VALIDA["nombre"]


# ---------- Editar (botón "Editar", solo mientras siga en revisión) ----------

def test_administrador_puede_editar_una_solicitud_en_revision(client):
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]

    r = client.put(f"/solicitudes-alta-medico/{folio}", headers=admin, json={
        **SOLICITUD_VALIDA,
        "nombre": "Dra. Nombre Corregido",
        "calle_num_ext_int": "Av. Corregida 500, Int. 9",
        "telefono": "33 1111 2222",
    })
    assert r.status_code == 200, r.text

    solicitud = client.get("/solicitudes-alta-medico", headers=admin).json()[0]
    assert solicitud["nombre"] == "Dra. Nombre Corregido"
    assert solicitud["calle_num_ext_int"] == "Av. Corregida 500, Int. 9"
    assert solicitud["telefono"] == "33 1111 2222"
    assert solicitud["estatus"] == "En revisión", "editar no debe resolverla"
    assert solicitud["folio"] == folio, "el folio no cambia"


def test_administrador_puede_editar_una_con_cedula_no_localizada(client):
    """El caso que de verdad hay que poder corregir: la cédula profesional no
    se localizó en el registro, así que hay que arreglar el dato y volver a
    intentar (decisión del usuario, 2026-10-19)."""
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    marcar = client.post(
        f"/solicitudes-alta-medico/{folio}/estatus", headers=admin,
        json={"estatus": "Cédula no localizada"},
    )
    assert marcar.status_code == 200, marcar.text

    r = client.put(f"/solicitudes-alta-medico/{folio}", headers=admin, json={
        **SOLICITUD_VALIDA, "cedula": "87654321",
    })
    assert r.status_code == 200, r.text

    solicitud = client.get("/solicitudes-alta-medico", headers=admin).json()[0]
    assert solicitud["cedula"] == "87654321"
    assert solicitud["estatus"] == "Cédula no localizada", "editar no cambia el estatus"


def test_marcar_cedula_no_localizada_no_da_de_alta_en_el_padron(client):
    """Solo "Aprobada" crea el proveedor; los demás estatus no tocan el padrón."""
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    client.post(
        f"/solicitudes-alta-medico/{folio}/estatus", headers=admin,
        json={"estatus": "Cédula no localizada"},
    )
    assert client.get("/directorio-proveedores", headers=admin).json() == []


def test_no_se_edita_una_solicitud_ya_aprobada(client):
    """Ya aprobada, el proveedor existe en el padrón con su clave y puede
    tener citas: editar aquí dejaría los dos lados en desacuerdo."""
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Aprobada"})

    r = client.put(
        f"/solicitudes-alta-medico/{folio}", headers=admin,
        json={**SOLICITUD_VALIDA, "nombre": "Dr. Ya No Se Puede"},
    )
    assert r.status_code == 409
    assert "aprobada" in r.json()["detail"].lower()
    assert client.get("/solicitudes-alta-medico", headers=admin).json()[0]["nombre"] == SOLICITUD_VALIDA["nombre"]


def test_no_se_edita_una_solicitud_ya_rechazada(client):
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Rechazada"})

    r = client.put(
        f"/solicitudes-alta-medico/{folio}", headers=admin,
        json={**SOLICITUD_VALIDA, "nombre": "Dr. Ya No Se Puede"},
    )
    assert r.status_code == 409
    assert "rechazada" in r.json()["detail"].lower()


def test_editar_folio_inexistente_regresa_404(client):
    r = client.put(
        "/solicitudes-alta-medico/ALT-00000", headers=_auth_headers_admin(client), json=SOLICITUD_VALIDA
    )
    assert r.status_code == 404


def test_editar_rechaza_cabina(client):
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    r = client.put(f"/solicitudes-alta-medico/{folio}", headers=_auth_headers_cabina(client), json=SOLICITUD_VALIDA)
    assert r.status_code == 403


# ---------- Flujo completo ----------

def test_crear_solicitud_queda_en_revision_y_con_folio(client):
    r = client.post("/solicitudes-alta-medico", headers=_auth_headers_proveedor(client), json=SOLICITUD_VALIDA)
    assert r.status_code == 200
    folio = r.json()["folio"]
    assert folio.startswith("ALT-")

    listado = client.get("/solicitudes-alta-medico", headers=_auth_headers_admin(client)).json()
    assert len(listado) == 1
    assert listado[0]["folio"] == folio
    assert listado[0]["estatus"] == "En revisión"
    assert listado[0]["nombre"] == SOLICITUD_VALIDA["nombre"]


def test_solicitud_aprobada_no_aparece_en_el_directorio_antes_de_aprobarse(client):
    headers = _auth_headers_proveedor(client)
    client.post("/solicitudes-alta-medico", headers=headers, json=SOLICITUD_VALIDA)
    directorio = client.get("/directorio-proveedores", headers=headers).json()
    assert directorio == []


def test_solicitud_aprobada_entra_al_padron_con_clave_propia_y_origen_portal(client):
    admin = _auth_headers_admin(client)
    r = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA)
    folio = r.json()["folio"]

    aprobar = client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Aprobada"})
    assert aprobar.status_code == 200
    clave = aprobar.json()["clave_proveedor"]
    assert clave.startswith("PRV-")

    directorio = client.get("/directorio-proveedores", headers=_auth_headers_proveedor(client)).json()
    assert len(directorio) == 1
    assert directorio[0]["origen"] == "portal"       # "Fuera de la red"
    assert directorio[0]["clave"] == clave
    assert directorio[0]["folio_alta"] == folio
    assert directorio[0]["cl_proveedor"] is None
    assert directorio[0]["nombre"] == SOLICITUD_VALIDA["nombre"]
    # El domicilio de consulta llega desglosado, igual que en la carga.
    assert directorio[0]["calle_num_ext_int"] == SOLICITUD_VALIDA["calle_num_ext_int"]
    assert directorio[0]["colonia"] == SOLICITUD_VALIDA["colonia"]
    assert directorio[0]["codigo_postal"] == SOLICITUD_VALIDA["codigo_postal"]


def test_aprobar_dos_veces_no_duplica_al_medico_en_el_padron(client):
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    primera = client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Aprobada"})
    segunda = client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Aprobada"})

    assert primera.json()["clave_proveedor"] == segunda.json()["clave_proveedor"]
    assert len(client.get("/directorio-proveedores", headers=admin).json()) == 1


def test_rechazar_solicitud_no_la_vuelve_visible_en_el_directorio(client):
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Rechazada"})

    directorio = client.get("/directorio-proveedores", headers=_auth_headers_proveedor(client)).json()
    assert directorio == []

    listado = client.get("/solicitudes-alta-medico", headers=admin).json()
    assert listado[0]["estatus"] == "Rechazada"


def _cargar_excel(client, headers, filas: list[dict]):
    from io import BytesIO

    from openpyxl import Workbook

    from app.services.proveedores_excel_service import COLUMNAS

    wb = Workbook()
    ws = wb.active
    for col_idx, (encabezado, _clave, _ancho, obligatorio) in enumerate(COLUMNAS, start=1):
        ws.cell(row=1, column=col_idx, value=encabezado + (" *" if obligatorio else ""))
    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, (_encabezado, clave, _ancho, _obligatorio) in enumerate(COLUMNAS, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=fila.get(clave))
    buffer = BytesIO()
    wb.save(buffer)
    return client.post(
        "/proveedores-catalogo/cargar", headers=headers,
        files={"archivo": (
            "catalogo.xlsx", buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
    )


def test_una_carga_de_excel_no_borra_a_los_medicos_dados_de_alta_en_el_portal(client):
    """La carga es incremental: traer el padrón del sistema central no debe
    tirar a los médicos que se dieron de alta aquí."""
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Aprobada"})

    _cargar_excel(client, admin, [
        {"nombre": "Dr. Del Catálogo", "especialidad": "Pediatría", "cedula": "999999"},
    ])

    directorio = client.get("/directorio-proveedores", headers=_auth_headers_proveedor(client)).json()
    assert {d["origen"] for d in directorio} == {"catalogo", "portal"}
    assert len(directorio) == 2


def test_cuando_el_medico_del_portal_llega_en_la_carga_conserva_su_clave_y_pasa_a_la_red(client):
    """Lo pidió el usuario (2026-10-19): el alta del portal "se quita" cuando
    el médico ya llega en la carga -- pero conservando su identificador, que
    es lo que amarra las citas ya generadas. En vez de borrarlo y crear otro,
    el MISMO renglón pasa de origen 'portal' ("Fuera de la red") a 'catalogo'
    ("Dentro de la red")."""
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    aprobar = client.post(
        f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Aprobada"}
    )
    clave = aprobar.json()["clave_proveedor"]

    # Una cita agendada ANTES de que llegue en la carga.
    from datetime import date
    cita = client.post("/directorio-proveedores/citas", headers=admin, json={
        "clave_proveedor": clave, "fecha": date.today().isoformat(), "hora": "09:00",
    })
    assert cita.status_code == 200, cita.text

    # Ahora el mismo médico (misma cédula y especialidad) llega en el Excel.
    r = _cargar_excel(client, admin, [{
        "nombre": SOLICITUD_VALIDA["nombre"], "especialidad": SOLICITUD_VALIDA["especialidad"],
        "cedula": SOLICITUD_VALIDA["cedula"], "cl_proveedor": 777, "costo_consulta": "1200",
    }])
    assert r.json() == {"ok": True, "total": 1, "nuevos": 0, "actualizados": 0, "absorbidos": 1}

    directorio = client.get("/directorio-proveedores", headers=admin).json()
    assert len(directorio) == 1, "no debe duplicarse"
    assert directorio[0]["clave"] == clave, "la clave se conserva"
    assert directorio[0]["origen"] == "catalogo", "ya quedó dentro de la red"
    assert directorio[0]["cl_proveedor"] == 777
    assert directorio[0]["costo_consulta"] == "1200"

    # Y la cita sigue apuntando al mismo médico.
    citas = client.get("/directorio-proveedores/citas", headers=admin).json()
    assert len(citas) == 1
    assert citas[0]["clave_proveedor"] == clave
    assert citas[0]["proveedor_nombre"] == SOLICITUD_VALIDA["nombre"]


def test_estatus_invalido_se_rechaza(client):
    admin = _auth_headers_admin(client)
    folio = client.post("/solicitudes-alta-medico", headers=admin, json=SOLICITUD_VALIDA).json()["folio"]
    r = client.post(f"/solicitudes-alta-medico/{folio}/estatus", headers=admin, json={"estatus": "Lo que sea"})
    assert r.status_code == 400


def test_estatus_de_folio_inexistente_regresa_404(client):
    admin = _auth_headers_admin(client)
    r = client.post("/solicitudes-alta-medico/ALT-00000/estatus", headers=admin, json={"estatus": "Aprobada"})
    assert r.status_code == 404
