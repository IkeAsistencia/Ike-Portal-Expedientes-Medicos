"""
Pruebas de /proveedores-catalogo (carga del catálogo de proveedores del
Directorio de Proveedores, exclusiva de Administrador) -- NO requiere SQL
Server real: todo vive en SQLite local.
"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services.proveedores_catalogo_excel_service import COLUMNAS, generar_plantilla_excel


@pytest.fixture(autouse=True)
def _sin_sql_server(monkeypatch):
    import app.main as main_module

    monkeypatch.setattr(main_module, "calentar_pool", lambda: None)


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_AUTH_MODE", "sql")
    monkeypatch.setenv("DB_SERVER", "servidor-fake")
    monkeypatch.setenv("DB_NAME", "bd-fake")
    monkeypatch.setenv("DB_USER", "usuario-fake")
    monkeypatch.setenv("DB_PASSWORD", "password-fake")
    monkeypatch.setenv("LOCAL_DB_PATH", str(tmp_path / "test_local.db"))
    monkeypatch.setenv("JWT_SECRET_KEY", "clave-de-prueba")

    from app.config import get_settings

    get_settings.cache_clear()

    with TestClient(app) as c:
        yield c


def _auth_headers_admin(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCADMTEST", "Admin de Prueba", accesos_repo_module.PERFIL_ADMINISTRADOR)
    accesos_repo_module.crear_password("RFCADMTEST", "Admin2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCADMTEST", "password": "Admin2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_proveedor(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCPROTEST", "Proveedor de Prueba", accesos_repo_module.PERFIL_PROVEEDOR)
    accesos_repo_module.crear_password("RFCPROTEST", "Proveedor2024")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCPROTEST", "password": "Proveedor2024"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _auth_headers_cabina(client):
    import app.repositories.accesos_repo as accesos_repo_module

    accesos_repo_module.alta_acceso("RFCCABTEST", "Cabina de Prueba", accesos_repo_module.PERFIL_CABINA)
    accesos_repo_module.crear_password("RFCCABTEST", "Cabina2024xx")
    r = client.post("/auth/rfc/login", json={"rfc": "RFCCABTEST", "password": "Cabina2024xx"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _xlsx_de_filas(filas: list[dict]) -> bytes:
    """Arma un .xlsx con los mismos encabezados de la plantilla real, para
    no depender de generar_plantilla_excel() salvo en la prueba que la
    prueba explícitamente a ella."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for col_idx, (encabezado, _clave, _ancho, obligatorio) in enumerate(COLUMNAS, start=1):
        ws.cell(row=1, column=col_idx, value=encabezado + (" *" if obligatorio else ""))
    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, (_encabezado, clave, _ancho, _obligatorio) in enumerate(COLUMNAS, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=fila.get(clave))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


FILA_VALIDA = {
    "cl_proveedor": 10234, "nombre": "Dra. Laura Méndez Ruiz", "especialidad": "Cardiología",
    "cedula": "8453219", "rfc": "MERL850214H23", "estado": "Jalisco", "municipio": "Guadalajara",
    "direccion": "Av. Vallarta 1450", "consultorio": "Consultorio 302", "telefono": "33 1234 5678",
    "email": "consultorio@correo.mx", "horario": "Lun-Vie 09:00-14:00", "costo_consulta": "800",
    "convenio": "GNP", "teleconsulta": "No",
}


# ---------- Acceso exclusivo de Administrador ----------

def test_plantilla_requiere_administrador(client):
    r = client.get("/proveedores-catalogo/plantilla", headers=_auth_headers_proveedor(client))
    assert r.status_code == 403


def test_cargar_requiere_administrador(client):
    archivo = _xlsx_de_filas([FILA_VALIDA])
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_cabina(client),
        files={"archivo": ("catalogo.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 403


def test_resumen_requiere_administrador(client):
    r = client.get("/proveedores-catalogo", headers=_auth_headers_proveedor(client))
    assert r.status_code == 403


# ---------- Plantilla ----------

def test_plantilla_se_descarga_y_abre_como_excel_valido(client):
    r = client.get("/proveedores-catalogo/plantilla", headers=_auth_headers_admin(client))
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = load_workbook(BytesIO(r.content))
    encabezados = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    assert encabezados[0] == "Clave Proveedor (Core) *"
    assert encabezados[1] == "Nombre completo *"


def test_la_plantilla_generada_se_puede_leer_de_vuelta(client):
    """La plantilla que se descarga (con su fila de ejemplo) debe pasar la
    misma validación que un archivo real -- si no, el ejemplo que le damos
    al área estaría mal."""
    from app.services.proveedores_catalogo_excel_service import leer_catalogo_excel

    registros = leer_catalogo_excel(generar_plantilla_excel())
    assert len(registros) == 1
    assert registros[0]["cl_proveedor"] == 10234


# ---------- Carga: reemplazo completo ----------

def test_cargar_catalogo_lo_deja_disponible_en_el_resumen(client):
    headers = _auth_headers_admin(client)
    archivo = _xlsx_de_filas([FILA_VALIDA, {**FILA_VALIDA, "cl_proveedor": 55555, "nombre": "Dr. Otro Médico"}])
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=headers,
        files={"archivo": ("catalogo.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200
    assert r.json() == {"ok": True, "total": 2}

    resumen = client.get("/proveedores-catalogo", headers=headers).json()
    assert resumen["total"] == 2
    assert resumen["ultima_carga"]["rfc"] == "RFCADMTEST"
    assert resumen["ultima_carga"]["total_filas"] == 2


def test_segunda_carga_reemplaza_la_anterior_por_completo(client):
    """Decisión del usuario (2026-10-16): cada carga sustituye TODO el
    catálogo, no hace upsert -- si el segundo archivo trae menos filas, el
    catálogo final debe quedar con esas menos filas, no acumular."""
    headers = _auth_headers_admin(client)
    primero = _xlsx_de_filas([FILA_VALIDA, {**FILA_VALIDA, "cl_proveedor": 2}])
    client.post(
        "/proveedores-catalogo/cargar", headers=headers,
        files={"archivo": ("a.xlsx", primero, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    segundo = _xlsx_de_filas([{**FILA_VALIDA, "cl_proveedor": 99}])
    r = client.post(
        "/proveedores-catalogo/cargar", headers=headers,
        files={"archivo": ("b.xlsx", segundo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.json()["total"] == 1

    resumen = client.get("/proveedores-catalogo", headers=headers).json()
    assert resumen["total"] == 1


# ---------- Validaciones del archivo ----------

def test_rechaza_archivo_que_no_es_xlsx(client):
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.txt", b"no es un excel", "text/plain")},
    )
    assert r.status_code == 400


def test_rechaza_xlsx_falso_con_content_type_mentiroso(client):
    """El Content-Type lo manda el cliente -- se valida también la firma
    real del archivo (un .xlsx es un ZIP, empieza con PK\\x03\\x04)."""
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": (
            "catalogo.xlsx", b"esto no es un zip valido",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
    )
    assert r.status_code == 400


def test_rechaza_archivo_sin_la_columna_obligatoria(client):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Nombre completo *")  # falta "Clave Proveedor (Core) *"
    ws.cell(row=2, column=1, value="Dr. Alguien")
    buffer = BytesIO()
    wb.save(buffer)

    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 400
    assert "Clave Proveedor" in r.json()["detail"]


def test_rechaza_fila_sin_nombre(client):
    archivo = _xlsx_de_filas([{**FILA_VALIDA, "nombre": None}])
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 400
    assert "Nombre completo" in r.json()["detail"]


def test_rechaza_cl_proveedor_no_numerico(client):
    archivo = _xlsx_de_filas([{**FILA_VALIDA, "cl_proveedor": "ABC"}])
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 400
    assert "número entero" in r.json()["detail"]


def test_ignora_renglones_vacios_al_final(client):
    archivo = _xlsx_de_filas([FILA_VALIDA, {}, {}])
    r = client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200
    assert r.json()["total"] == 1


def test_teleconsulta_si_no_se_interpreta_a_booleano(client):
    import app.repositories.proveedores_catalogo_repo as repo

    archivo = _xlsx_de_filas([
        {**FILA_VALIDA, "cl_proveedor": 1, "teleconsulta": "Sí"},
        {**FILA_VALIDA, "cl_proveedor": 2, "teleconsulta": "No"},
    ])
    client.post(
        "/proveedores-catalogo/cargar",
        headers=_auth_headers_admin(client),
        files={"archivo": ("catalogo.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    catalogo = {r["cl_proveedor"]: r for r in repo.listar_catalogo()}
    assert catalogo[1]["teleconsulta"] == 1
    assert catalogo[2]["teleconsulta"] == 0
